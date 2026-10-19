"""0.4: check_tracker / fix_tracker, the same-company rule, reply dates and reply_times,
the follow-up calendar, and matching e-mails to applications. On the `new` template,
which has every column (furthest stage, reply date, follow-up log)."""

import json
from datetime import date, datetime

import anyio
import openpyxl
import pytest
from mcp import Client

from tracker_mcp import server
from tracker_mcp.config import load_config
from tracker_mcp.insights import to_ics
from tracker_mcp.mail import classify
from tracker_mcp.scaffold import new_tracker
from tracker_mcp.tracker import Tracker, TrackerError, company_key

TODAY = date(2026, 10, 19)   # a Monday
COLS = {"company": "B", "role": "C", "applied_on": "H", "channel": "I", "status": "J", "reached": "K",
        "response": "N", "replied_on": "O", "followup_on": "P", "followup_log": "R", "notes": "S"}


def build(tmp_path, rows):
    """A `new` tracker with these rows: dicts of field -> value."""
    path, _ = new_tracker(tmp_path / "jobs.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Applications"]
    for i, row in enumerate(rows, start=2):
        ws[f"A{i}"] = i - 1
        for field, value in row.items():
            ws[f"{COLS[field]}{i}"] = value
    wb.save(path)
    return Tracker(path, config=load_config(path), today=lambda: TODAY)


def test_the_template_has_a_reply_date_and_init_finds_one(tmp_path):
    t = build(tmp_path, [])
    assert t.cfg.columns["replied_on"] == "Replied on" and t.cfg.is_date("replied_on")
    for header in ("Дата ответа", "Reply date", "Date of response"):
        wb = openpyxl.Workbook()
        wb.active.append(["Company", "Role", "Status", "Applied on", "Response", header])
        p = tmp_path / f"{company_key(header)}.xlsx"
        wb.save(p)
        from tracker_mcp.scaffold import init
        _, g, _ = init(p)
        assert g.columns["replied_on"][1] == header and g.columns["response"][1] == "Response"


# ---------------------------------------------------------------- check / fix

def messy(tmp_path):
    return build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": "2026-09-10", "status": "Applied",
         "followup_on": datetime(2026, 10, 25), "followup_log": "2026-09-10: applied"},          # Sunday
        {"company": "Globex", "role": "Head of Eng", "applied_on": 46280, "status": "Rejected",
         "response": "Auto-confirmation 10.09; rejection 24.09 from the recruiter",
         "followup_on": datetime(2026, 10, 30)},                                              # closed, planned
        {"company": "Initech", "role": "Director"},                                            # no status
        {"company": "Hooli", "role": "VP", "applied_on": datetime(2026, 9, 1), "status": "interview",
         "response": "Thank you for applying; call with the CTO 2026-09-15"},                  # case; reached
        {"company": "Umbrella", "role": "EM", "applied_on": "2026-08", "status": "Applied"},     # bad date, no FU
        {"company": "Acme Inc.", "role": "EM", "status": "Not applied"},                        # duplicate
    ])


def test_check_finds_what_hand_edits_leave_behind(tmp_path):
    t = messy(tmp_path)
    res = t.check_tracker()
    got = {(i["kind"], i["row"]): i for i in res["issues"]}
    assert got[("date_as_text", 2)]["fix"] == "2026-09-10"
    assert got[("weekend_followup", 2)]["fix"] == "2026-10-26"
    assert got[("date_as_text", 3)]["fix"] == "2026-09-15"                       # an Excel serial number
    assert got[("closed_with_followup", 3)]["fix"] == "—"
    assert got[("no_reply_date", 3)]["fix"] == "2026-09-24"                      # not the confirmation's date
    assert "rejection 24.09" in got[("no_reply_date", 3)]["note"]
    assert got[("no_status", 4)]["fix"] == "Not applied"
    assert got[("unknown_status", 5)]["fix"] == "Interview"                      # the dropdown's spelling
    assert got[("reached_behind", 5)]["fix"] == "Interview"
    assert got[("no_reply_date", 5)]["fix"] == "2026-09-15"
    assert "fix" not in got[("bad_date", 6)] and ("no_followup", 6) in got
    assert ("duplicate", 7) in got and "row 2" in got[("duplicate", 7)]["note"]
    assert res["fixable"] == sum(1 for i in res["issues"] if "fix" in i)


def test_fix_writes_the_fixable_ones_in_one_undoable_step(tmp_path):
    t = messy(tmp_path)
    before = t.version()
    with pytest.raises(TrackerError, match="Unknown kind"):
        t.fix_tracker(kinds=["typos"], dry_run=True)
    only_dates = t.fix_tracker(kinds=["date_as_text"], dry_run=True)
    assert {c["field"] for c in only_dates["changes"]} == {"applied_on"}
    res = t.fix_tracker()
    assert res["row"] == [2, 3, 4, 5]
    # the text date on row 2 and its Sunday follow-up: one value per cell, the later check wins
    fields = [(c["row"], c["field"]) for c in res["changes"]]
    assert len(fields) == len(set(fields))
    ws = openpyxl.load_workbook(t.path)["Applications"]
    assert ws["H2"].value == datetime(2026, 9, 10) and ws["P2"].value == datetime(2026, 10, 26)
    assert ws["J4"].value == "Not applied" and ws["O3"].value == datetime(2026, 9, 24)
    left = {i["kind"] for i in t.check_tracker()["issues"]}
    assert left == {"bad_date", "no_followup", "duplicate"}
    assert t.check_tracker()["fixable"] == 0
    t.undo_last_write()
    assert t.version() == before


# ---------------------------------------------------------------- rules

def test_another_application_to_the_same_company_is_a_warning(tmp_path):
    t = build(tmp_path, [
        {"company": "Swap (Swap Commerce)", "role": "Director", "applied_on": datetime(2026, 10, 7),
         "status": "Applied"},
        {"company": "SWAP Inc.", "role": "VP Eng", "status": "Not applied"},
        {"company": "Old Co", "role": "EM", "applied_on": datetime(2026, 8, 1), "status": "Rejected"},
    ])
    plan = t.mark_applied(3, dry_run=True)
    assert len(plan["warnings"]) == 1 and "Row 2" in plan["warnings"][0] and "12 days" in plan["warnings"][0]
    assert not t.add_application("Old Co", "Director", dry_run=True).get("warnings")   # 79 days ago
    assert t.add_application("Swap", "CTO", dry_run=True)["warnings"]
    text = server.describe_plan("mark_applied", plan)
    assert "⚠ Row 2" in text
    off = Tracker(t.path, config=t.cfg.model_copy(update={"rules": t.cfg.rules.model_copy(
        update={"same_company_days": 0})}), today=lambda: TODAY)
    assert not off.mark_applied(3, dry_run=True).get("warnings")


def test_the_first_reply_fills_the_reply_date(tmp_path):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Applied"}])
    with pytest.raises(TrackerError, match="future"):
        t.set_status(2, "Screening", on=date(2026, 10, 20), dry_run=True)
    t.set_status(2, "Screening", on=date(2026, 9, 12))
    t.set_status(2, "Rejected")
    rec = t.get_application(row=2)["applications"][0]
    assert rec["replied_on"] == "2026-09-12" and rec["reached"] == "Screening"
    (tmp_path / "b").mkdir()
    t2 = build(tmp_path / "b", [{"company": "Acme", "role": "EM", "status": "Applied"}])
    t2.set_status(2, "Withdrawn")                       # not a reply
    assert t2.get_application(row=2)["applications"][0]["replied_on"] is None


# ---------------------------------------------------------------- reply times

def test_reply_times_and_the_silence_table(tmp_path):
    rows = []
    for i, (applied, replied, status, channel) in enumerate([
            (date(2026, 8, 1), date(2026, 8, 4), "Rejected", "LinkedIn"),
            (date(2026, 8, 1), date(2026, 8, 11), "Screening", "Referral"),
            (date(2026, 8, 3), date(2026, 8, 30), "Rejected", "LinkedIn"),
            (date(2026, 8, 5), None, "Applied", "LinkedIn"),
            (date(2026, 8, 6), None, "No response", "LinkedIn"),
            (date(2026, 10, 15), None, "Applied", "Referral"),
            (date(2026, 8, 7), None, "Rejected", "LinkedIn")]):   # a reply with no date
        rows.append({"company": f"Co{i}", "role": "EM", "applied_on": datetime(*applied.timetuple()[:3]),
                     "status": status, "channel": channel,
                     **({"replied_on": datetime(*replied.timetuple()[:3])} if replied else {})})
    t = build(tmp_path, rows)
    res = t.reply_times(by="channel")
    assert res["total"]["replies_with_dates"] == 3 and res["total"]["median_days"] == 10
    assert res["replies_without_a_date"] == 1 and res["source"] == "the replied_on column"
    silence = {s["after_days_silent"]: s for s in res["silence"]}
    # after 7 silent days: Co1 (10 d) and Co2 (27 d) replied later; Co3, Co4 never did
    assert (silence[7]["applications"], silence[7]["replied_later"]) == (4, 2)
    assert (silence[30]["applications"], silence[30]["replied_later"]) == (2, 0)
    groups = {g["group"]: g for g in res["groups"]}
    assert groups["LinkedIn"]["replies_with_dates"] == 2 and groups["Referral"]["applied"] == 2


def test_reply_times_without_a_column_uses_the_servers_own_records(tmp_path):
    path, _ = new_tracker(tmp_path / "jobs.xlsx")
    cfg = load_config(path)
    cols = {k: v for k, v in cfg.columns.items() if k != "replied_on"}
    t = Tracker(path, config=cfg.model_copy(update={"columns": cols}), today=lambda: TODAY)
    t.add_application("Acme", "EM")
    t.mark_applied(2, applied_on=date(2026, 10, 1))
    t.set_status(2, "Interview")
    res = t.reply_times()
    assert res["total"]["replies_with_dates"] == 1 and res["total"]["median_days"] == 18
    assert "recorded a reply" in res["source"]


# ---------------------------------------------------------------- calendar

def test_calendar_file(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme, Inc.", "role": "EM; Platform", "applied_on": datetime(2026, 9, 1), "status": "Applied",
         "followup_on": datetime(2026, 10, 12), "followup_log": "x" * 300},                     # overdue
        {"company": "Globex", "role": "EM", "applied_on": datetime(2026, 10, 9), "status": "Applied",
         "followup_on": datetime(2026, 10, 21)},
        {"company": "Initech", "role": "EM", "status": "Rejected", "followup_on": datetime(2026, 10, 22)},
        {"company": "Hooli", "role": "EM", "status": "Applied", "followup_on": datetime(2027, 3, 1)},  # too far
    ])
    res = t.export_calendar()
    assert res["file"].endswith("jobs.followups.ics") and res["events"] == 2 and res["overdue_on_today"] == 1
    text = (tmp_path / "jobs.followups.ics").read_bytes().decode()
    assert text.startswith("BEGIN:VCALENDAR\r\n") and "\r\nEND:VCALENDAR\r\n" in text
    assert all(len(line.encode()) <= 75 for line in text.split("\r\n"))
    unfolded = text.replace("\r\n ", "")
    assert "SUMMARY:Overdue Follow-up: Acme\\, Inc. — EM\\; Platform" in unfolded
    assert "DTSTART;VALUE=DATE:20261019" in unfolded and "DTSTART;VALUE=DATE:20261021" in unfolded
    uid = [line for line in unfolded.split("\r\n") if line.startswith("UID:")]
    t.export_calendar()
    again = (tmp_path / "jobs.followups.ics").read_bytes().decode().replace("\r\n ", "")
    assert [line for line in again.split("\r\n") if line.startswith("UID:")] == uid   # stable ids
    with pytest.raises(TrackerError, match=".ics"):
        t.export_calendar(tmp_path / "out.txt")


def test_ics_folding_keeps_multibyte_characters_whole():
    ics = to_ics([{"uid": "u@x", "date": "2026-10-19", "title": "Фоллоу-ап: " + "Компания " * 20,
                   "description": "строка\nвторая"}])
    for line in ics.split("\r\n"):
        assert len(line.encode()) <= 75
        line.encode().decode()   # no character cut in half


# ---------------------------------------------------------------- e-mail

@pytest.mark.parametrize("subject,snippet,kind", [
    ("Thank you for applying to Acme", "If there is a match we will schedule a call", "confirmation"),
    ("Update on your application", "Unfortunately we decided to move forward with other candidates", "rejection"),
    ("Interview invitation: Acme", "We would love to chat", "interview"),
    ("Thanks for applying to Globex", "We would like to invite you to a call with our CTO", "interview"),
    ("Your offer", "We are pleased to offer you the role", "offer"),
    ("Ваш отклик", "К сожалению, мы не готовы продолжить", "rejection"),
    ("Newsletter", "Ten tips for your CV", "other"),
])
def test_classify(subject, snippet, kind):
    assert classify({"subject": subject, "snippet": snippet}) == kind


def test_match_emails_to_rows(tmp_path):
    t = build(tmp_path, [
        {"company": "AIOS", "role": "Head of AI", "applied_on": datetime(2026, 9, 10), "status": "Applied"},
        {"company": "Supabase", "role": "EM", "applied_on": datetime(2026, 9, 11), "status": "Applied"},
        {"company": "Supabase", "role": "Director", "applied_on": datetime(2026, 9, 18), "status": "Applied"},
        {"company": "Canonical", "role": "Director", "applied_on": datetime(2026, 9, 11), "status": "Applied"},
        {"company": "bunny.net", "role": "EM", "status": "Not applied"},
    ])
    res = t.match_emails([
        {"from": "Ashby <no-reply@ashbyhq.com>", "subject": "Ivan Rykov <> Aios Medical",
         "snippet": "Unfortunately we will not be moving forward", "date": "Thu, 24 Sep 2026 07:15:00 +0000"},
        {"from": "no-reply@ashbyhq.com", "subject": "Thank you for applying to Supabase!", "date": "2026-09-11"},
        {"from": "Varshi Gupta <varshi.gupta@canonical.com>", "subject": "Engineering Director",
         "snippet": "We would like to invite you to the written interview stage", "date": "2026-10-07"},
        {"from": "Daria <daria@bunny.net>", "subject": "Your application to bunny.net",
         "snippet": "Thank you for applying!", "date": "2026-09-17"},
        {"from": "LinkedIn <jobs-noreply@linkedin.com>", "subject": "Your application was sent to Hooli"},
    ])
    m = res["messages"]
    assert m[0]["kind"] == "rejection" and [r["row"] for r in m[0]["rows"]] == [2]
    assert m[0]["suggestion"] == {"tool": "set_status", "args": {"row": 2, "status": "Rejected", "on": "2026-09-24"}}
    assert len(m[1]["rows"]) == 2 and "ask the user" in m[1]["suggestion"]["note"]
    assert m[2]["suggestion"]["args"] == {"row": 5, "status": "Screening", "on": "2026-10-07"}
    assert m[3]["suggestion"] == {"tool": "mark_applied", "args": {"row": 6, "applied_on": "2026-09-17"}}
    assert m[4]["rows"] == [] and m[4]["company_in_mail"] == "Hooli" and "add_application" in m[4]["suggestion"]["note"]
    assert res["matched"] == 4


# ---------------------------------------------------------------- over MCP

def test_new_tools_over_mcp(tmp_path):
    server._tracker = messy(tmp_path)

    async def go():
        async with Client(server.mcp) as c:
            tools = {x.name: x for x in (await c.list_tools()).tools}
            for name in ("check_tracker", "reply_times", "match_emails"):
                assert tools[name].annotations.read_only_hint is True
            assert tools["fix_tracker"].annotations.destructive_hint is True
            assert tools["export_followups_calendar"].annotations.destructive_hint is False
            assert "on" in tools["set_status"].input_schema["properties"]
            res = await c.call_tool("fix_tracker", {})
            assert res.is_error and "confirmation" in res.content[0].text
            preview = await c.call_tool("fix_tracker", {"kinds": ["no_status"], "dry_run": True})
            token = preview.structured_content["confirm"]
            res = await c.call_tool("fix_tracker", {"kinds": ["no_status"], "confirm": token})
            assert not res.is_error and res.structured_content["row"] == [4]
            res = await c.call_tool("match_emails", {"messages": [
                {"from": "x@globex.com", "subject": "Unfortunately", "date": "2026-10-01"}]})
            assert res.structured_content["messages"][0]["rows"][0]["company"] == "Globex"
            prompts = {p.name for p in (await c.list_prompts()).prompts}
            assert {"weekly_review", "check_mail"} <= prompts
            text = (await c.get_prompt("check_mail", {"days": "3"})).messages[0].content.text
            assert "last 3 days" in text and "match_emails" in text

    anyio.run(go)


def test_calendar_cli(tmp_path, capsys, monkeypatch):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 10, 9), "status": "Applied",
                          "followup_on": datetime(2026, 12, 1)}])
    monkeypatch.delenv("TRACKER_PATH", raising=False)
    out = tmp_path / "f.ics"
    assert server.calendar_cli(["--tracker", str(t.path), "--out", str(out), "--days", "365"]) == 0
    assert "1 follow-up" in capsys.readouterr().out and "Acme" in out.read_text(encoding="utf-8")
    assert server.calendar_cli(["--tracker", str(t.path), "--out", str(tmp_path / "x.txt")]) == 2


def test_example_config_has_the_new_keys():
    from tracker_mcp.scaffold import example_config
    text = example_config()
    assert "same_company_days" in text and "silence_from" in text and "replied_on" in text
    json.dumps(text)

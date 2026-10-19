"""CSV and Google Sheets trackers, and the fixes from the fresh-install review."""

import csv
import os
from datetime import date, datetime

import openpyxl
import pytest

from tracker_mcp import gsheets
from tracker_mcp.config import load_config
from tracker_mcp.scaffold import init, new_tracker
from tracker_mcp.stores import ConflictError
from tracker_mcp.tracker import Tracker, TrackerError

from .fake_sheets import FakeSheets
from .test_tracker import RU_CONFIG, TODAY, make_tracker, ru_tracker

URL = "https://docs.google.com/spreadsheets/d/abc123XYZ/edit#gid=0"


# ---------------------------------------------------------------- CSV

def test_csv_tracker_end_to_end(tmp_path):
    book, cfg = new_tracker(tmp_path / "jobs.csv")
    t = Tracker(book, config=load_config(book), today=lambda: TODAY)
    assert t.describe()["kind"] == "csv" and t.describe()["data_rows"] == "none yet"
    t.add_application("Acme", "EM", extra={"Found via": "LinkedIn"})
    t.mark_applied(2, channel="Ashby (direct form)")
    t.set_status(2, "Interview", response="call on Monday")
    head, row = list(csv.reader(book.open(encoding="utf-8")))[:2]
    got = dict(zip(head, row, strict=True))
    assert (got["#"], got["Company"], got["Found via"]) == ("1", "Acme", "LinkedIn")
    assert got["Applied on"] == "2026-10-18"                                    # dates as yyyy-mm-dd
    assert got["Status"] == got["Furthest stage"] == "Interview"
    t.undo_last_write()
    assert dict(zip(head, list(csv.reader(book.open(encoding="utf-8")))[1], strict=True))["Status"] == "Applied"
    assert t.funnel()["total"]["applied"] == 1


def test_csv_keeps_its_delimiter_and_bom(tmp_path):
    path = tmp_path / "export.csv"
    path.write_bytes("﻿Company;Role;Status;Date Applied\nAcme;EM;Applied;09/02/2026\n".encode())
    _, g, text = init(path)
    assert g.date_formats == ["%m/%d/%Y"]                       # US dates found and configured
    t = Tracker(path, config=load_config(path), today=lambda: TODAY)
    assert t.list_applications()["applications"][0]["applied_on"] == "2026-09-02"
    t.set_status(2, "Rejected")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and b"Acme;EM;Rejected;09/02/2026" in raw


# ---------------------------------------------------------------- Google Sheets

@pytest.fixture
def sheet(tmp_path, monkeypatch):
    monkeypatch.setenv("TRACKER_MCP_HOME", str(tmp_path / "home"))
    fake = FakeSheets({
        "Jobs": [["Company", "Role", "Status", "Date Applied", "Follow-up date", "Follow-up log", "Notes"],
                 ["Acme", "EM", "Applied", datetime(2026, 9, 1), datetime(2026, 9, 11), None, None],
                 ["Globex", "Head of Eng", "Wishlist", None, None, None, None]],
        "Stats": [["Applied", "=COUNTIF(Jobs!C2:C3,\"Applied\")"]],
    }, dropdowns={"Jobs": {3: ["Wishlist", "Applied", "Interview", "Offer", "Rejected", "Ghosted"]}})
    monkeypatch.setattr(gsheets, "http_transport", lambda auth: fake)
    return fake


def test_google_sheet_init_and_writes(sheet, tmp_path):
    target, g, _ = init(URL)
    assert target == tmp_path / "home" / "abc123XYZ.tracker-mcp.toml"
    assert g.statuses_from == "dropdown" and g.summary_sheet == "Stats"
    t = Tracker(URL, config=load_config(URL), today=lambda: TODAY)
    d = t.describe()
    assert d["kind"] == "gsheet" and d["statuses_allowed"][0] == "Wishlist"

    res = t.mark_applied(3, channel="Referral")
    assert sheet.value("Jobs", 3, 3) == "Applied"
    assert sheet.value("Jobs", 3, 4) == datetime(2026, 10, 18)              # written as a real date
    assert sheet.is_date("Jobs", 3, 4)
    assert (t.backup_dir / res["backup"].split("/")[-1]).exists()

    t.add_application("Initech", "VP", notes="=IMPORTXML(\"http://evil\")")   # text stays text
    assert sheet.sheets["Jobs"]["kinds"][(4, 7)] == "stringValue"
    with pytest.raises(TrackerError, match="Allowed"):
        t.set_status(2, "Hired")


def test_google_sheet_conflict_and_undo(sheet):
    init(URL)
    t = Tracker(URL, config=load_config(URL), today=lambda: TODAY)
    v = t.version()
    sheet.set("Jobs", 2, 7, "edited in the browser")
    with pytest.raises(ConflictError):
        t.set_status(2, "Interview", expected_version=v)
    t.set_status(2, "Interview")
    assert sheet.value("Jobs", 2, 3) == "Interview"
    t.undo_last_write()
    assert sheet.value("Jobs", 2, 3) == "Applied"


def test_google_sheet_grows_the_grid(sheet):
    init(URL)
    sheet.sheets["Jobs"]["rows"] = 3          # a grid with no spare rows
    t = Tracker(URL, config=load_config(URL), today=lambda: TODAY)
    t.add_application("Hooli", "EM")
    assert sheet.value("Jobs", 4, 1) == "Hooli"
    assert any("appendDimension" in r for r in sheet.requests)


def test_google_sign_in_is_explained(tmp_path, monkeypatch):
    monkeypatch.setenv("TRACKER_MCP_HOME", str(tmp_path))
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    with pytest.raises(TrackerError, match="tracker-mcp auth google"):
        gsheets.GoogleAuth().header()


# ---------------------------------------------------------------- review fixes

def test_extra_cannot_set_managed_columns(tmp_path):
    t = ru_tracker(tmp_path / "t.xlsx")
    for key in ("Статус", "status", "Дата подачи", "Фоллоу-ап — статус"):
        with pytest.raises(TrackerError, match="set by the server"):
            t.add_application("X", "Y", extra={key: "whatever"}, dry_run=True)


def test_list_rejects_an_unknown_status(tmp_path):
    t = ru_tracker(tmp_path / "t.xlsx")
    with pytest.raises(TrackerError, match="Unknown status"):
        t.list_applications(status="Interviewing")
    assert t.list_applications(status="нет ответа")["count"] == 0       # an alias is fine
    assert t.list_applications(status="подана")["count"] == 3


def test_a_locked_file_is_a_conflict_not_a_crash(tmp_path, monkeypatch):
    import tracker_mcp.stores as stores
    t = ru_tracker(tmp_path / "t.xlsx")
    before = t.version()

    def locked(src, dst):
        raise PermissionError(13, "The process cannot access the file")

    monkeypatch.setattr(stores.os, "replace", locked)
    with pytest.raises(ConflictError, match="holding it"):
        t.set_status(7, "Отказ")
    assert t.version() == before


def test_describe_tells_what_is_missing(tmp_path):
    path = tmp_path / "min.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Company", "Job Title", "Status", "Location"])
    wb.active.append(["Acme", "EM", "Applied", "Remote"])
    wb.save(path)
    _, g, _ = init(path)
    assert any("no follow-up date column" in c for c in g.checks)
    d = Tracker(path, config=load_config(path)).describe()
    assert d["other_headers"] == ["Location"]
    assert set(d["unavailable_tools"]) == {"followups_due", "log_followup", "reschedule_followup",
                                           "stale_applications", "close_stale"}
    f = Tracker(path, config=load_config(path)).funnel("Location")
    assert f["groups"][0]["group"] == "Remote"


def test_ambiguous_dates_are_flagged(tmp_path):
    path = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Company", "Role", "Status", "Applied on"])
    wb.active.append(["Acme", "EM", "Applied", "03/04/2026"])
    wb.active.append(["Globex", "EM", "Applied", "not sure"])
    wb.save(path)
    _, g, _ = init(path)
    assert g.date_formats == ["%m/%d/%Y"]
    assert any("month-first or day-first" in c for c in g.checks)
    assert any("«not sure»" in c for c in g.checks)


def test_cli_takes_the_tracker_as_an_option(tmp_path):
    import subprocess
    import sys
    book = make_tracker(tmp_path / "t.xlsx")
    cfg = tmp_path / "ru.toml"
    cfg.write_text(RU_CONFIG.read_text(encoding="utf-8"), encoding="utf-8")
    env = {k: v for k, v in os.environ.items() if not k.startswith("TRACKER_")}
    out = subprocess.run([sys.executable, "-m", "tracker_mcp.server", "--tracker", str(book), "--config",
                          str(cfg), "--check"], capture_output=True, text=True, encoding="utf-8", env=env)
    assert out.returncode == 0, out.stderr
    assert '"kind": "xlsx"' in out.stdout
    ex = subprocess.run([sys.executable, "-m", "tracker_mcp.server", "example"], capture_output=True,
                        text=True, encoding="utf-8")
    assert ex.returncode == 0 and "[statuses.stages]" in ex.stdout


def test_dates_written_from_extra_are_dates(tmp_path):
    path = tmp_path / "d.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Company", "Role", "Status", "Interview on"])
    wb.save(path)
    cfg = tmp_path / "d.tracker-mcp.toml"
    cfg.write_text('sheet = "Sheet"\ndate_columns = ["interview_on"]\n[columns]\ncompany = "Company"\n'
                   'role = "Role"\nstatus = "Status"\ninterview_on = "Interview on"\n', encoding="utf-8")
    t = Tracker(path, config=load_config(path))
    with pytest.raises(TrackerError, match="not a date"):
        t.add_application("Acme", "EM", extra={"interview_on": "next week"}, dry_run=True)
    t.add_application("Acme", "EM", extra={"Interview on": "2026-10-21"})
    assert openpyxl.load_workbook(path).active["D2"].value == datetime(2026, 10, 21)
    assert date(2026, 10, 21)


def test_first_start_writes_the_config(tmp_path, monkeypatch):
    """Installed as an extension, nobody runs `init`: the server writes the guess itself."""
    import subprocess
    import sys

    from tracker_mcp import server
    path = tmp_path / "jobs.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Jobs"
    wb.active.append(["Company", "Role", "Status", "Applied on"])
    wb.active.append(["Acme", "EM", "Applied", "2026-10-01"])
    wb.save(path)
    env = {k: v for k, v in os.environ.items() if not k.startswith("TRACKER_")}
    out = subprocess.run([sys.executable, "-m", "tracker_mcp.server", "--tracker", str(path), "--check"],
                         capture_output=True, text=True, encoding="utf-8", env=env)
    assert out.returncode == 2 and "tracker-mcp init" in out.stderr   # --check never writes
    assert not (tmp_path / "jobs.tracker-mcp.toml").exists()

    monkeypatch.setenv("TRACKER_PATH", str(path))
    monkeypatch.delenv("TRACKER_CONFIG", raising=False)
    monkeypatch.setattr(server, "_tracker", None)
    monkeypatch.setattr(server, "_watch", None)
    monkeypatch.setattr(server, "_generated", None)
    layout = server._describe()
    assert (tmp_path / "jobs.tracker-mcp.toml").exists()
    assert layout["config_generated"]["file"] == str(tmp_path / "jobs.tracker-mcp.toml")
    assert layout["sheet"] == "Jobs" and layout["columns"]["company"] == "A: Company"
    assert server.call("list_applications")["applications"][0]["company"] == "Acme"


def test_text_that_looks_like_a_formula_stays_text(tmp_path):
    for suffix in (".xlsx", ".csv"):
        book, _ = new_tracker(tmp_path / f"{suffix[1:]}{suffix}")
        t = Tracker(book, config=load_config(book))
        payload = '=HYPERLINK("http://evil.example","click")'
        t.add_application("Acme", "EM", notes=payload)
        t.add_application("Globex", "EM", notes="- referral from Ann")
        apps = {a["company"]: a for a in t.get_application("EM")["applications"]}
        if suffix == ".xlsx":
            wb = openpyxl.load_workbook(book)
            ws = wb[t.cfg.sheet]
            cells = [c for row in ws.iter_rows() for c in row if c.value == payload]
            assert cells and all(c.data_type == "s" for c in cells)
            assert apps["Acme"]["notes"] == payload
        else:
            assert apps["Acme"]["notes"] == "'" + payload
            assert "'=HYPERLINK" in book.read_text(encoding="utf-8")
        assert apps["Globex"]["notes"] == "- referral from Ann"

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

from tracker_mcp.config import load_config
from tracker_mcp.tracker import ConflictError, Tracker, TrackerError, parse_date

RU_CONFIG = Path(__file__).parent / "fixtures" / "ru.tracker-mcp.toml"

HEADERS = ["№", "Компания", "Роль", "Источник", "География — проверено", "Деньги (опубликовано)",
           "Ссылка на подачу", "Папка", "Что подготовлено", "Гейт — спросить первым", "Подался",
           "Дата подачи", "Канал подачи", "Ответ", "Статус", "Заметки", "Фоллоу-ап 1 — дата",
           "Фоллоу-ап — кому / канал", "Фоллоу-ап — статус", "Разбор — главная причина",
           "Деньги / локация / уровень (0–2)"]
STATUSES = "Не подавался,Подана,Ответа нет,Отказ,Скрининг,Интервью,Тех.этап,Финал,Оффер,Отозвал"
TODAY = date(2026, 10, 18)  # a Sunday


def make_tracker(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Подачи"
    ws["A1"] = "Трекер подач"
    for c, h in enumerate(HEADERS, 1):
        ws.cell(5, c, h)
    rows = [
        # №, company, role, applied, applied_on, status, followup_on, followup_log
        (1, "Initech", "Director of Engineering", "Да", "2026-09-10", "Отказ", None, "Отменён"),
        ("2", "Globex", "Senior EM, Platform", "Да", datetime(2026, 10, 7), "Подана",
         datetime(2026, 10, 19), "07.10: подано"),
        (3, "Acme", "Engineering Manager, Agents", "Да", 46313, "Подана", "46323", None),
        (4, "Umbrella", "Engineering Manager, Integrations", "Нет", None, "Не подавался", None, None),
        (5, "Hooli", "EM (Billing)", "Да", datetime(2026, 10, 18), "Подана", None, None),
    ]
    for i, (num, co, role, applied, on, st, fu, log) in enumerate(rows, start=6):
        for col, value in {1: num, 2: co, 3: role, 11: applied, 12: on, 15: st,
                           17: fu, 19: log, 21: "2 / 0 / 1"}.items():
            ws.cell(i, col, value)
    for col in ("K", "O"):
        dv = DataValidation(type="list", formula1=f'"{"Да,Нет" if col == "K" else STATUSES}"')
        dv.add(f"{col}6:{col}10")
        ws.add_data_validation(dv)
    ws.auto_filter.ref = "A5:U10"
    s = wb.create_sheet("Сводка")
    s["B4"] = "=COUNTA(Подачи!B6:B12)"
    s["B5"] = '=COUNTIF(Подачи!K6:K12,"Да")'
    wb.save(path)
    return path


@pytest.fixture
def t(tmp_path):
    return ru_tracker(tmp_path / "tracker.xlsx")


def ru_tracker(path):
    """The Russian-language layout of my real tracker, from tests/fixtures."""
    return Tracker(make_tracker(path), config=load_config(path, RU_CONFIG), today=lambda: TODAY)


def test_parse_date_handles_every_format_in_the_file():
    assert parse_date("2026-09-10") == date(2026, 9, 10)
    assert parse_date(datetime(2026, 10, 7, 0, 0)) == date(2026, 10, 7)
    assert parse_date(46313) == date(2026, 10, 18)       # Excel serial
    assert parse_date("46323") == date(2026, 10, 28)     # serial stored as text
    assert parse_date("18.10.2026") == date(2026, 10, 18)
    assert parse_date("—") is None and parse_date("2026-08") is None


def test_reads(t):
    assert t.list_applications(status="Подана")["count"] == 3
    assert t.list_applications(applied_since=date(2026, 10, 1))["count"] == 3
    acme = t.get_application("acme")["applications"][0]
    assert acme["row"] == 8 and acme["applied_on"] == "2026-10-18"
    assert acme["other"] == {"Деньги / локация / уровень (0–2)": "2 / 0 / 1"}
    assert acme["money"] is None  # «Деньги» maps to F, not to the analysis column U

    due = t.followups_due(within_days=7)
    assert [d["company"] for d in due["due"]] == ["Globex"]
    assert [d["company"] for d in due["applied_without_followup_date"]] == ["Hooli"]

    s = t.summary(since=date(2026, 10, 1), until=date(2026, 10, 31))
    assert s["applied"] == 4 and s["applied_by_month"] == {"2026-09": 1, "2026-10": 3}
    assert s["period"]["applied"] == 3


def test_mark_applied_sets_followup_on_a_weekday(t):
    res = t.mark_applied(9, channel="Greenhouse")
    fields = {c["field"]: c["new"] for c in res["changes"]}
    assert fields["applied"] == "Да" and fields["status"] == "Подана"
    assert fields["applied_on"] == "2026-10-18"
    assert fields["followup_on"] == "2026-10-28"         # +10 days, a Wednesday
    assert t.get_application(row=9)["applications"][0]["status"] == "Подана"
    with pytest.raises(TrackerError, match="already marked"):
        t.mark_applied(9)


def test_dry_run_writes_nothing(t):
    v = t.version()
    res = t.set_status(7, "Отказ", response="template", dry_run=True)
    assert res["dry_run"] and res["changes"] and t.version() == v
    assert not t.backup_dir.exists()


def test_closing_status_cancels_followup_and_backs_up(t):
    res = t.set_status(7, "отказ", response="Шаблон, без фидбэка")
    got = t.get_application(row=7)["applications"][0]
    assert got["status"] == "Отказ" and got["followup_on"] == "—"
    assert got["response"] == "18.10: Шаблон, без фидбэка"
    assert "снят" in got["followup_log"]
    assert (t.backup_dir / res["backup"].split("/")[-1]).exists()
    assert t.recent_writes()["writes"][0]["tool"] == "set_status"


def test_status_must_come_from_the_dropdown(t):
    assert t.set_status(8, "Нет ответа", dry_run=True)["changes"][0]["new"] == "Ответа нет"
    with pytest.raises(TrackerError, match="Allowed"):
        t.set_status(8, "Ждём")
    with pytest.raises(TrackerError, match="mark_applied"):
        t.set_status(9, "Подана")


def test_conflicts(t):
    old = t.version()
    t.log_followup(7, "InMail Kaushik")
    with pytest.raises(ConflictError, match="changed since"):
        t.set_status(7, "Скрининг", expected_version=old)
    lock = t.path.parent / f"~${t.path.name}"
    lock.write_text("")
    with pytest.raises(ConflictError, match="open in Excel"):
        t.set_status(7, "Скрининг")
    assert t.set_status(7, "Скрининг", dry_run=True)["dry_run"]  # preview still works
    lock.unlink()


def test_add_grows_dropdowns_filter_and_summary(t):
    with pytest.raises(TrackerError, match="already has"):
        t.add_application("acme", "Engineering Manager, Agents")
    for i in range(3):
        res = t.add_application(f"Co{i}", "EM")
    assert res["row"] == 13
    wb = openpyxl.load_workbook(t.path)
    ws = wb["Подачи"]
    assert ws["A13"].value == 8 and ws["O13"].value == "Не подавался"
    assert {str(dv.sqref) for dv in ws.data_validations.dataValidation} == {"K6:K13", "O6:O13"}
    assert ws.auto_filter.ref == "A5:U13"
    assert wb["Сводка"]["B4"].value == "=COUNTA(Подачи!B6:B113)"


def test_undo_only_when_nothing_changed_since(t):
    before = t.version()
    t.log_followup(7, "FU-1 sent")
    t.undo_last_write()
    assert t.version() == before
    with pytest.raises(TrackerError):
        t.undo_last_write()
    t.log_followup(7, "FU-1 sent")
    openpyxl.load_workbook(t.path).save(t.path)  # someone else saves the file
    with pytest.raises(ConflictError, match="changed after"):
        t.undo_last_write()


def test_bad_row(t):
    with pytest.raises(TrackerError, match="not an application row"):
        t.set_status(5, "Отказ")


def test_close_stale_in_one_batch(tmp_path):
    path = make_tracker(tmp_path / "tracker.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Подачи"]
    ws["L7"], ws["Q7"] = datetime(2026, 9, 1), datetime(2026, 9, 11)   # Globex: silent since September
    ws["L9"], ws["O9"], ws["K9"] = "2026-09-02", "Ждём ответа", "Да"     # Umbrella: waiting, no date
    ws["Q10"] = datetime(2026, 10, 25)                                  # Hooli: follow-up still ahead
    wb.save(path)
    t = Tracker(path, config=load_config(path, RU_CONFIG), today=lambda: TODAY)

    stale = t.stale_applications(30)
    assert [a["company"] for a in stale["applications"]] == ["Globex", "Umbrella"]
    assert stale["would_set_status"] == "Ответа нет"

    with pytest.raises(TrackerError, match="not stale"):
        t.close_stale(30, rows=[8], dry_run=True)
    before = t.version()
    res = t.close_stale(30)
    assert res["row"] == [7, 9] and set(res["labels"]) == {7, 9}
    got = {r: t.get_application(row=r)["applications"][0] for r in (7, 9)}
    assert all(g["status"] == "Ответа нет" and g["followup_on"] == "—" for g in got.values())
    assert got[7]["followup_log"].endswith("18.10: закрыто: 30 дн. без ответа")
    assert len(list(t.backup_dir.glob("*.xlsx"))) == 1          # one backup for the batch
    t.undo_last_write()
    assert t.version() == before                                # one undo for the batch

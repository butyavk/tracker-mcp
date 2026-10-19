"""0.4 edge cases found in review: what check/fix must leave alone, which replies count,
what resets the silence clock, and the e-mail cases that matched the wrong row."""

from datetime import date, datetime

import openpyxl
import pytest

from tracker_mcp import server
from tracker_mcp.config import load_config
from tracker_mcp.doctor import reply_dates_from_audit
from tracker_mcp.insights import followup_events, to_ics
from tracker_mcp.mail import _domain_company, _msg_date, classify
from tracker_mcp.scaffold import new_tracker
from tracker_mcp.tracker import Tracker, TrackerError

from .test_insights import TODAY, build


def kinds(t):
    return {(i["kind"], i["row"]): i for i in t.check_tracker()["issues"]}


def without_reply_column(tmp_path):
    path, _ = new_tracker(tmp_path / "jobs.xlsx")
    cfg = load_config(path)
    cols = {k: v for k, v in cfg.columns.items() if k != "replied_on"}
    return Tracker(path, config=cfg.model_copy(update={"columns": cols}), today=lambda: TODAY)


# ---------------------------------------------------------------- check / fix

def test_a_date_with_more_text_is_reported_but_never_rewritten(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": "2026-09-10 via referral from Ann", "status": "Applied",
         "followup_on": datetime(2026, 10, 27)},
        {"company": "Globex", "role": "EM", "applied_on": "2026-09-12", "status": "Applied",
         "followup_on": datetime(2026, 10, 28)},
    ])
    got = kinds(t)
    assert "fix" not in got[("date_with_text", 2)] and "via referral" in got[("date_with_text", 2)]["note"]
    plan = t.fix_tracker(dry_run=True)
    assert {(c["row"], c["field"]) for c in plan["changes"]} == {(3, "applied_on")}
    # the preview shows what the cell held, as text, and the date it becomes
    assert "applied_on: «2026-09-12» (text) → 2026-09-12 (a date)" in server.describe_plan("fix_tracker", plan)
    t.fix_tracker()
    ws = openpyxl.load_workbook(t.path)["Applications"]
    assert ws["H2"].value == "2026-09-10 via referral from Ann" and ws["H3"].value == datetime(2026, 9, 12)


def test_a_fixed_date_takes_the_format_of_the_columns_other_dates(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Rejected"},
        {"company": "Globex", "role": "EM", "applied_on": "11.09.2026", "status": "Rejected"},
    ])
    wb = openpyxl.load_workbook(t.path)
    wb["Applications"]["H2"].number_format = "dd.mm.yyyy"
    wb.save(t.path)
    t.fix_tracker(kinds=["date_as_text"])
    cell = openpyxl.load_workbook(t.path)["Applications"]["H3"]
    assert cell.value == datetime(2026, 9, 11) and cell.number_format == "dd.mm.yyyy"


def test_an_alias_is_what_the_later_checks_see(tmp_path):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Ghosted",
                          "followup_on": datetime(2026, 10, 25)}])   # a Sunday ahead
    st = t.cfg.statuses.model_copy(update={"aliases": {"Ghosted": "No response"}})
    t = Tracker(t.path, config=t.cfg.model_copy(update={"statuses": st}), today=lambda: TODAY)
    got = kinds(t)
    assert got[("unknown_status", 2)]["fix"] == "No response"
    # closed, so the follow-up goes; not moved to Monday as if the row were still open
    assert ("closed_with_followup", 2) in got and ("weekend_followup", 2) not in got


def test_no_status_follows_the_submit_date(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1)},
        {"company": "Globex", "role": "EM"},
    ])
    got = kinds(t)
    assert got[("no_status", 2)]["fix"] == "Applied" and "submit date" in got[("no_status", 2)]["note"]
    assert got[("no_status", 3)]["fix"] == "Not applied"


def test_merged_rows_are_skipped(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": "2026-09-01", "status": "Applied",
         "followup_on": datetime(2026, 10, 27)},
        {"company": "— Paused —"},
    ])
    wb = openpyxl.load_workbook(t.path)
    wb["Applications"].merge_cells("B3:R3")
    wb.save(t.path)
    assert {row for _, row in kinds(t)} == {2}
    t.fix_tracker()
    assert openpyxl.load_workbook(t.path)["Applications"]["H2"].value == datetime(2026, 9, 1)


@pytest.mark.parametrize("response,expected", [
    ("Auto-confirmation 2026-09-02; rejection 2026-09-20", date(2026, 9, 20)),
    ("Автоподтверждение 02.09; отказ 20.09", date(2026, 9, 20)),
    ("Thank you for applying 02.09 | call with the CTO 15.09", date(2026, 9, 15)),
    ("Applied 01.09 via the site; screening invite 2026-09-05", date(2026, 9, 5)),
    ("We have received your application (02.09)", None),
])
def test_dates_of_confirmations_are_not_reply_dates(tmp_path, response, expected):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1),
                          "status": "Rejected", "response": response}])
    issue = kinds(t)[("no_reply_date", 2)]
    if expected:
        assert issue["fix"] == expected.isoformat()
    else:
        assert "fix" not in issue


# ---------------------------------------------------------------- reply dates

def test_the_audit_counts_only_writes_that_stand_on_rows_that_did_not_move(tmp_path):
    t = without_reply_column(tmp_path)
    t.add_application("Acme", "EM")
    t.add_application("Globex", "EM")
    for row in (2, 3):
        t.mark_applied(row, applied_on=date(2026, 10, 1))
    t.set_status(2, "Interview")
    t.undo_last_write()
    t.set_status(3, "Screening")
    labels = {2: "Acme — EM", 3: "Globex — EM"}
    assert reply_dates_from_audit(t, labels) == {3: TODAY}
    # after a sort, row 3 holds another application: its old reply is not this one's
    assert reply_dates_from_audit(t, {2: "Acme — EM", 3: "Initech — EM"}) == {}


def test_only_the_first_reply_sets_the_reply_date(tmp_path):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Screening"}])
    t.set_status(2, "Interview", on=date(2026, 9, 20))    # already replied before: the date stays empty
    assert t.get_application(row=2)["applications"][0]["replied_on"] is None
    ws_path = t.path
    wb = openpyxl.load_workbook(ws_path)
    wb["Applications"]["O2"] = datetime(2026, 9, 5)
    wb.save(ws_path)
    t.set_status(2, "Rejected", on=date(2026, 10, 1))
    assert t.get_application(row=2)["applications"][0]["replied_on"] == "2026-09-05"


def test_reply_times_rejects_an_unknown_grouping_and_skips_withdrawn_rows(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": datetime(2026, 8, 1), "status": "Withdrawn"},
        {"company": "Globex", "role": "EM", "applied_on": datetime(2026, 8, 1), "status": "No response"},
        {"company": "Initech", "role": "EM", "applied_on": datetime(2026, 8, 1), "status": "Rejected",
         "replied_on": datetime(2026, 8, 3)},
    ])
    with pytest.raises(TrackerError, match="Cannot group by"):
        t.reply_times(by="colour")
    res = t.reply_times()
    assert res["total"]["applied"] == 2 and res["total"]["replies_with_dates"] == 1


# ---------------------------------------------------------------- silence

def test_moving_or_dropping_a_follow_up_does_not_reset_the_silence(tmp_path):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Applied"}])
    t.reschedule_followup(2, None, "waiting for LinkedIn credits")
    rec = t.stale_applications(days=30)["applications"]
    assert rec and rec[0]["last_touch"] == "2026-09-01" and rec[0]["days_silent"] == 48
    (tmp_path / "b").mkdir()
    t = build(tmp_path / "b", [{"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1),
                                "status": "Applied", "followup_log": "2026-09-01: applied | 2026-10-01: InMail "
                                "to the hiring manager | 2026-10-15: follow-up moved to 2026-10-16: on holiday"}])
    assert not t.stale_applications(days=30)["applications"]
    assert t.stale_applications(days=18)["applications"][0]["last_touch"] == "2026-10-01"


def test_mark_applied_logs_the_submit_day_not_today(tmp_path):
    t = build(tmp_path, [{"company": "Acme", "role": "EM", "status": "Not applied"}])
    t.mark_applied(2, applied_on=date(2026, 10, 2))
    log = t.get_application(row=2)["applications"][0]["followup_log"]
    assert log.startswith("2026-10-02") and "2026-10-19" not in log


# ---------------------------------------------------------------- e-mail

@pytest.mark.parametrize("subject,snippet,kind", [
    ("Your interview with Acme", "Unfortunately we need to move your interview to Thursday", "interview"),
    ("Acme: next steps", "We'd love to offer you an interview slot next week", "interview"),
    ("Application received", "We will schedule a call if your profile matches", "confirmation"),
    ("Your application", "We have received your application and will schedule interviews soon", "confirmation"),
    ("Update", "We regret to inform you that we will not schedule an interview", "rejection"),
    ("Приглашение на собеседование", "Добрый день! Приглашаем на созвон", "interview"),
])
def test_classify_edges(subject, snippet, kind):
    assert classify({"subject": subject, "snippet": snippet}) == kind


def test_sender_domains_and_dates():
    assert _domain_company("Jane <jane@acme.co.uk>") == "acme"
    assert _domain_company("jobs@careers.globex.com") == "globex"
    assert _domain_company("no-reply@us.greenhouse-mail.io") is None
    assert _msg_date("1790000000000", TODAY) == datetime.fromtimestamp(1790000000).date()
    assert _msg_date("2030-01-01", TODAY) == TODAY          # never in the future
    assert _msg_date("not a date", TODAY) is None


def test_a_persons_name_does_not_match_a_company_by_its_start(tmp_path):
    t = build(tmp_path, [
        {"company": "Deel", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Applied"},
        {"company": "Swap", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Applied"},
        {"company": "Globex", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "No response"},
    ])
    res = t.match_emails([
        {"from": "Deeleep Kumar <deeleep@gmail.com>", "subject": "Coffee?"},
        {"from": "Swapnil Rao <swapnil.rao@outlook.com>", "subject": "Hello"},
        {"from": "Ann <ann@globex.com>", "subject": "Your application", "snippet": "We would like to invite you "
         "to a call", "date": "2026-10-10"},
        "not a message",
    ])
    m = res["messages"]
    assert m[0]["rows"] == [] and m[1]["rows"] == []
    # closed for no response, but a late reply is still news
    assert m[2]["suggestion"]["args"] == {"row": 4, "status": "Screening", "on": "2026-10-10"}
    assert "error" in m[3]


# ---------------------------------------------------------------- calendar

def test_calendar_ids_and_carriage_returns(tmp_path):
    t = build(tmp_path, [
        {"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 1), "status": "Applied",
         "followup_on": datetime(2026, 10, 21), "followup_log": "line one\r\nline two\rthree"},
        {"company": "Acme", "role": "EM", "applied_on": datetime(2026, 9, 2), "status": "Applied",
         "followup_on": datetime(2026, 10, 22)},
    ])
    events, _ = followup_events(t)
    assert len({e["uid"] for e in events}) == 2           # the same company and role, two ids
    ics = to_ics(events)
    assert "\r" not in ics.replace("\r\n", "")
    assert "line one\\nline two\\nthree" in ics.replace("\r\n ", "")

"""The same server on a different tracker: English headers, other sheet names,
header in row 1, no running number, no yes/no column, no status dropdown."""

import json
import os
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from tracker_mcp.config import ConfigError, load_config
from tracker_mcp.tracker import Tracker, TrackerError

TODAY = date(2026, 10, 16)  # a Friday
CONFIG = """
sheet = "My Applications"
summary_sheet = "Stats"

[columns]
company = "Company"
role = "Position"
status = "Stage"
link = "URL"
applied_on = "Applied"
followup_on = "Follow up"
followup_log = "Log"
salary = "Salary"

[statuses]
not_applied = "Draft"
applied = "Sent"
closed = ["Rejected", "Offer"]
fallback = ["Draft", "Sent", "Interview", "Rejected", "Offer"]
aliases = { "ghosted" = "Rejected" }

[followups]
first_after_days = 7
next_after_days = 7
weekend_to_monday = false

[messages]
applied = "sent"
applied_with_followup = "sent; nudge on {date}"
followup_cancelled = "nudge {date} dropped ({status})"
"""


@pytest.fixture
def en(tmp_path):
    path = tmp_path / "jobs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "My Applications"
    ws.append(["Company", "Position", "Stage", "URL", "Applied", "Follow up", "Log", "Salary", "Comment"])
    ws.append(["Acme", "EM", "Sent", "https://acme/1", datetime(2026, 10, 1), datetime(2026, 10, 15), None, None, "hi"])
    ws.append(["Globex", "Head of Eng", "Draft", None, None, None, None, "200k", None])
    stats = wb.create_sheet("Stats")
    stats["A1"] = "=COUNTA('My Applications'!A2:A3)"
    wb.save(path)
    (tmp_path / "jobs.tracker-mcp.toml").write_text(CONFIG, encoding="utf-8")
    return Tracker(path, config=load_config(path), today=lambda: TODAY)


def test_reads_with_another_layout(en):
    d = en.describe()
    assert d["header_row"] == 1 and d["columns"]["status"] == "C: Stage"
    assert d["custom_fields"] == ["salary"]
    assert d["statuses_allowed"] == ["Draft", "Sent", "Interview", "Rejected", "Offer"]  # fallback

    globex = en.get_application("globex")["applications"][0]
    assert globex["salary"] == "200k" and globex["row"] == 3
    acme = en.get_application(row=2)["applications"][0]
    assert acme["other"] == {"Comment": "hi"}

    s = en.summary()
    assert s["applied"] == 1 and s["not_applied"] == 1  # "applied" = status is not Draft
    due = en.followups_due(7)["due"]
    assert [x["company"] for x in due] == ["Acme"] and due[0]["overdue"]


def test_writes_follow_the_config(en):
    res = en.mark_applied(3)
    new = {c["field"]: c["new"] for c in res["changes"]}
    assert new["status"] == "Sent" and new["applied_on"] == "2026-10-16"
    assert new["followup_on"] == "2026-10-23"           # +7, no weekend shift configured
    assert new["followup_log"] == "2026-10-16: sent; nudge on 2026-10-23"
    with pytest.raises(TrackerError, match="already marked"):
        en.mark_applied(3)

    res = en.set_status(3, "ghosted")
    new = {c["field"]: c["new"] for c in res["changes"]}
    assert new["status"] == "Rejected" and new["followup_on"] == "—"
    assert new["followup_log"].endswith("nudge 2026-10-23 dropped (Rejected)")

    res = en.add_application("Initech", "VP Eng", extra={"salary": "250k", "Comment": "via friend"})
    assert res["row"] == 4
    wb = openpyxl.load_workbook(en.path)
    row = [c.value for c in wb["My Applications"][4]]
    assert row[:3] == ["Initech", "VP Eng", "Draft"] and row[7:] == ["250k", "via friend"]
    assert wb["Stats"]["A1"].value == "=COUNTA('My Applications'!A2:A104)"  # quoted sheet name

    with pytest.raises(TrackerError, match="Unknown column"):
        en.add_application("X", "Y", extra={"Nope": "1"}, dry_run=True)


def test_missing_column_disables_only_its_tools(tmp_path):
    path = tmp_path / "min.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active.append(["Company", "Role", "Status"])
    wb.active.append(["Acme", "EM", "Applied"])
    wb.save(path)
    cfg = tmp_path / "min.toml"
    cfg.write_text('sheet = "S"\n[columns]\ncompany = "Company"\nrole = "Role"\nstatus = "Status"\n')
    t = Tracker(path, config=load_config(path, cfg))
    assert t.summary()["applied"] == 1
    with pytest.raises(TrackerError, match="followup_on"):
        t.followups_due()


def test_config_errors(tmp_path):
    path = tmp_path / "x.xlsx"
    with pytest.raises(ConfigError, match="No config"):
        load_config(path)
    (tmp_path / "x.tracker-mcp.toml").write_text('sheet = "S"\n[columns]\ncompany = "C"\nrole = "R"\n')
    with pytest.raises(ConfigError, match="status"):
        load_config(path)
    (tmp_path / "x.tracker-mcp.toml").write_text(
        'sheet = "S"\nsheeet = "typo"\n[columns]\ncompany = "C"\nrole = "R"\nstatus = "S"\n')
    with pytest.raises(ConfigError, match="sheeet"):
        load_config(path)


def test_check_cli(en):
    out = subprocess.run([sys.executable, "-m", "tracker_mcp.server", "--check"], capture_output=True,
                         text=True, env={**os.environ, "TRACKER_PATH": str(en.path)})
    assert out.returncode == 0, out.stderr
    assert json.loads(out.stdout)["columns"]["role"] == "B: Position"
    bad = subprocess.run([sys.executable, "-m", "tracker_mcp.server", "--check"], capture_output=True,
                         text=True, env={**os.environ, "TRACKER_PATH": str(en.path),
                                         "TRACKER_CONFIG": str(Path(en.path).with_suffix(".none"))})
    assert bad.returncode == 2 and "No config" in bad.stderr


def test_example_config_is_valid():
    example = Path(__file__).parent.parent / "examples" / "tracker-mcp.example.toml"
    cfg = load_config("unused.xlsx", example)
    assert cfg.date_columns == ["applied_on", "followup_on"] and "date_columns" not in cfg.columns
    assert list(cfg.statuses.stages) == ["screening", "interview", "offer"]

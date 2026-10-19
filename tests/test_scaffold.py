"""Getting started (`init`, `new`) and the funnel."""

import subprocess
import sys
import tomllib
from datetime import date

import openpyxl
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

from tracker_mcp.config import load_config
from tracker_mcp.scaffold import guess, init, new_tracker
from tracker_mcp.tracker import Tracker, TrackerError, dropdown_values

from .test_tracker import RU_CONFIG, TODAY, make_tracker, ru_tracker


def test_init_rebuilds_my_hand_written_config(tmp_path):
    book = make_tracker(tmp_path / "t.xlsx")
    target, g, _ = init(book)
    assert target == tmp_path / "t.tracker-mcp.toml"
    hand = Tracker(book, config=load_config(book, RU_CONFIG)).describe()
    auto = Tracker(book, config=load_config(book)).describe()
    custom = {"source", "geography", "money", "folder", "prepared", "gate"}
    assert {k: v for k, v in hand["columns"].items() if k not in custom} == \
           {k: v for k, v in auto["columns"].items() if k != "source"}
    assert auto["columns"]["source"] == "D: Источник"
    assert (auto["header_row"], auto["data_rows"]) == (hand["header_row"], hand["data_rows"])

    h, a = load_config(book, RU_CONFIG), load_config(book)
    assert set(a.statuses.closed) == set(h.statuses.closed) - {"Закрыта"}  # not in this dropdown
    assert (a.statuses.not_applied, a.statuses.applied, a.statuses.no_response) == \
           (h.statuses.not_applied, h.statuses.applied, h.statuses.no_response)
    assert a.statuses.rejected == ["Отказ"] and a.statuses.stages == h.statuses.stages
    assert (a.values, a.messages, a.summary_sheet) == (h.values, h.messages, "Сводка")
    assert ("Гейт — спросить первым" in [hd for _, hd in g.unmapped]) and not g.checks


def test_init_on_a_sheet_without_dropdowns(tmp_path):
    path = tmp_path / "jobs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tab"
    ws.append(["Company", "Position", "Stage", "URL", "Applied", "Follow up", "Log", "Salary"])
    ws.append(["Acme", "EM", "Sent", "https://acme/1", "2026-10-01", None, None, "200k"])
    ws.append(["Globex", "Head of Eng", "Draft", None, None, None, None, None])
    wb.save(path)
    target, g, text = init(path)
    cfg = load_config(path)
    assert cfg.columns["applied_on"] == "Applied"          # a column of dates is the date of applying
    assert "applied" not in cfg.columns and cfg.columns["role"] == "Position"
    assert cfg.statuses.fallback[:2] == ["Sent", "Draft"] and cfg.statuses.applied == "Sent"
    assert {"Rejected", "No response", "Interview"} <= set(cfg.statuses.fallback)   # padded: no dropdown
    assert cfg.statuses.not_applied == "Draft" and cfg.statuses.no_response == "No response"
    assert any("added" in c for c in g.checks) and "# CHECK" in text
    assert "# salary = \"Salary\"" in text
    assert Tracker(path, config=cfg).summary()["applied"] == 1


def test_init_does_not_overwrite(tmp_path):
    book = make_tracker(tmp_path / "t.xlsx")
    init(book)
    with pytest.raises(TrackerError, match="already exists"):
        init(book)
    _, _, text = init(book, write=False)
    assert tomllib.loads(text)["sheet"] == "Подачи"
    with pytest.raises(TrackerError, match="No header row"):
        wb = openpyxl.Workbook()
        wb.active.append(["Name", "Phone"])
        wb.save(tmp_path / "contacts.xlsx")
        guess(tmp_path / "contacts.xlsx")


def test_dropdown_from_a_range_on_another_sheet(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Company", "Role", "Status"])
    lists = wb.create_sheet("Lists")
    for i, s in enumerate(["To apply", "Applied", "Ghosted", "Rejected"], 1):
        lists.cell(i, 1, s)
    dv = DataValidation(type="list", formula1="Lists!$A$1:$A$4")
    dv.add("C2:C100")
    ws.add_data_validation(dv)
    wb.save(tmp_path / "r.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "r.xlsx").active
    assert dropdown_values(ws, 3) == ["To apply", "Applied", "Ghosted", "Rejected"]
    g = guess(tmp_path / "r.xlsx")
    assert g.meaning["no_response"] == ["Ghosted"] and g.meaning["not_applied"] == ["To apply"]


def test_new_tracker_works_end_to_end(tmp_path):
    book, cfg_path = new_tracker(tmp_path / "jobs.xlsx")
    cfg = load_config(book)
    assert cfg.columns["reached"] == "Furthest stage" and cfg.summary_sheet == "Summary"
    t = Tracker(book, config=cfg, today=lambda: TODAY)
    for co in ("Acme", "Globex", "Initech"):
        t.add_application(co, "Engineering Manager", extra={"source": "LinkedIn"})
    ws = openpyxl.load_workbook(book)["Applications"]
    assert ws["A2"].value == 1 and not ws["B2"].font.b          # not styled like the header
    t.mark_applied(2, channel="Ashby (direct form)")
    t.mark_applied(3, channel="ashby")
    t.mark_applied(4, channel="Greenhouse")
    t.set_status(2, "Interview")
    t.set_status(2, "Rejected")
    t.set_status(3, "Screening")
    got = t.get_application(row=2)["applications"][0]
    assert (got["status"], got["reached"]) == ("Rejected", "Interview")  # the stage stays reached

    f = t.funnel()
    assert f["total"] == {**f["total"], "applied": 3, "replied": 2, "screening": 2, "interview": 1, "offer": 0}
    assert [(g["group"], g["applied"], g["interview"]) for g in f["groups"]] == [("Ashby", 2, 1),
                                                                               ("Greenhouse", 1, 0)]
    assert f["total"]["rates"]["replied"] == 66.7
    with pytest.raises(TrackerError, match="already exists"):
        new_tracker(book)


def test_funnel_on_my_layout(tmp_path):
    t = ru_tracker(tmp_path / "t.xlsx")
    t.set_status(7, "Тех.этап")
    t.set_status(7, "Отказ")          # no `reached` column here: counts as a reply only
    f = t.funnel("month")
    assert [g["group"] for g in f["groups"]] == ["2026-09", "2026-10"]
    assert f["total"]["applied"] == 4 and f["total"]["replied"] == 2 and f["total"]["interview"] == 0
    assert "current status only" in f["note"]
    assert t.funnel(since=date(2026, 10, 1))["total"]["applied"] == 3
    with pytest.raises(TrackerError, match="Cannot group by"):
        t.funnel("salary")


def test_cli(tmp_path):
    book = make_tracker(tmp_path / "t.xlsx")
    run = lambda *a: subprocess.run([sys.executable, "-m", "tracker_mcp.server", *a],  # noqa: E731
                                    capture_output=True, text=True, encoding="utf-8")
    out = run("init", str(book))
    assert out.returncode == 0, out.stderr
    assert "Компания" in out.stdout and "Next, check the mapping:  tracker-mcp --check --tracker" in out.stdout
    again = run("init", str(book))
    assert again.returncode == 2 and "already exists" in again.stderr
    assert "header_row = 5" in run("init", str(book), "--print").stdout
    created = run("new", str(tmp_path / "fresh.xlsx"))
    assert created.returncode == 0 and (tmp_path / "fresh.tracker-mcp.toml").exists()


def _book(path, headers, rows=(), dropdowns=None, sheet="Jobs"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for header, options in (dropdowns or {}).items():
        col = openpyxl.utils.get_column_letter(headers.index(header) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"')
        dv.add(f"{col}2:{col}100")
        ws.add_data_validation(dv)
    wb.save(path)
    return path


@pytest.mark.parametrize("word, kind", [
    ("Not interested", "withdrawn"), ("Interested", "not_applied"), ("Interview planned", "interview"),
    ("Offer Declined", "offer"), ("Declined offer", "offer"), ("Hired", "offer"),
    ("Waiting for offer", "final"), ("Phone Screen", "screening"), ("Onsite", "interview"),
    ("Coding challenge", "interview"), ("Ghosted", "no_response"), ("Not Selected", "rejected"),
    ("Application submitted", "applied"), ("Under review", "waiting"), ("Отказался", "withdrawn"),
    ("Отказ", "rejected"), ("Ждём ответа", "waiting"),
])
def test_status_words(word, kind):
    from tracker_mcp.scaffold import _classify
    meaning, unknown = _classify([word])
    assert meaning[kind] == [word] and not unknown


def test_rejected_after_a_stage_counts_for_that_stage():
    from tracker_mcp.scaffold import _classify
    meaning, _ = _classify(["Rejected after interview"])
    assert meaning["rejected"] == meaning["interview"] == ["Rejected after interview"]


def test_common_english_headers(tmp_path):
    headers = ["Job ID", "Company", "Job Title", "Source (LinkedIn/Indeed)", "Date Applied", "Status",
               "Response Date", "Contact", "Link", "Salary Range", "Notes", "Follow up"]
    g = guess(_book(tmp_path / "a.xlsx", headers, [["1", "Acme", "EM", "", "2026-10-01", "Applied"]]))
    got = {k: h for k, (_, h) in g.columns.items()}
    assert got["role"] == "Job Title" and got["link"] == "Link"
    assert got["source"] == "Source (LinkedIn/Indeed)" and got["applied_on"] == "Date Applied"
    assert got["followup_contact"] == "Contact" and got["followup_on"] == "Follow up"
    assert "response" not in got           # a date column must not take response text


def test_checkbox_applied_column(tmp_path):
    book = _book(tmp_path / "c.xlsx", ["Company", "Role", "Status", "Applied?"],
                 [["Acme", "EM", "Applied", True], ["Globex", "EM", "Draft", False]])
    _, g, _ = init(book)
    assert g.yes_no == ("TRUE", "FALSE")
    t = Tracker(book, config=load_config(book), today=lambda: TODAY)
    assert t.summary()["applied"] == 1
    with pytest.raises(TrackerError, match="already marked"):
        t.mark_applied(2)
    res = t.mark_applied(3)
    assert {c["field"]: c["new"] for c in res["changes"]}["applied"] is True
    t.add_application("Initech", "VP")
    assert openpyxl.load_workbook(book)["Jobs"]["D4"].value is False


def test_status_column_without_dropdown_or_values(tmp_path):
    book = _book(tmp_path / "e.xlsx", ["Company", "Role", "Status"], [["Acme", "EM", None]])
    _, g, text = init(book)
    assert any("default English statuses" in c for c in g.checks)
    assert load_config(book).statuses.fallback == ["Not applied", "Applied", "Screening", "Interview", "Offer",
                                                    "Rejected", "Withdrawn", "No response"]
    t = Tracker(book, config=load_config(book), today=lambda: TODAY)
    assert t.set_status(2, "Interview", dry_run=True)["changes"]


def test_a_status_with_a_line_break(tmp_path):
    book = _book(tmp_path / "n.xlsx", ["Company", "Role", "Status"],
                 [["Acme", "EM", "On hold\nuntil January"], ["Globex", "EM", "Applied"]])
    _, g, text = init(book)
    assert tomllib.loads(text) and any("On hold until January" in c for c in g.checks)


def test_reached_column_with_its_own_words(tmp_path):
    book = _book(tmp_path / "r.xlsx", ["Company", "Role", "Status", "Furthest stage", "Applied on"],
                 [["Acme", "EM", "Applied", None, "2026-10-01"], ["Globex", "EM", "Applied", None, "2026-10-01"]],
                 {"Status": ["To apply", "Applied", "Recruiter call", "Technical", "Offer", "Rejected", "Ghosted"],
                  "Furthest stage": ["Phone screen", "Final round", "Offer"]})
    init(book)
    cfg = load_config(book)
    assert "Final round" in cfg.statuses.stages["interview"] and "Phone screen" in cfg.statuses.stages["screening"]
    t = Tracker(book, config=cfg, today=lambda: TODAY)
    t.set_status(2, "Technical")
    assert t.get_application(row=2)["applications"][0]["reached"] == "Final round"   # the column's own word
    t.set_status(2, "Recruiter call")                                                  # never moves back
    assert t.get_application(row=2)["applications"][0]["reached"] == "Final round"
    ws = openpyxl.load_workbook(book)["Jobs"]
    ws["D3"] = "Coffee chat"
    ws.parent.save(book)
    res = t.set_status(3, "Technical")                    # an unknown word in reached is left alone
    assert [c["field"] for c in res["changes"]] == ["status"]
    assert t.funnel()["by"] == "month" and t.funnel()["total"]["interview"] == 2


def test_writes_refused_when_a_save_would_drop_parts_of_the_file(tmp_path, monkeypatch):
    import warnings

    import tracker_mcp.stores as tr
    real = tr.openpyxl.load_workbook

    def load(*a, **kw):
        warnings.warn("Data Validation extension is not supported and will be removed", UserWarning, stacklevel=2)
        return real(*a, **kw)

    t = ru_tracker(tmp_path / "t.xlsx")
    monkeypatch.setattr(tr.openpyxl, "load_workbook", load)
    assert t.describe()["writes_blocked"]
    assert "cannot be written" in t.set_status(7, "Отказ", dry_run=True)["warning"]
    with pytest.raises(TrackerError, match="Writes are off"):
        t.set_status(7, "Отказ")


def test_cli_errors_are_messages_not_tracebacks(tmp_path):
    book = make_tracker(tmp_path / "t.xlsx")
    run = lambda *a: subprocess.run([sys.executable, "-m", "tracker_mcp.server", *a],  # noqa: E731
                                    capture_output=True, text=True, encoding="utf-8")
    for args, msg in ((("init", str(book), "--sheet", "Nope"), "No sheet"),
                      (("new", str(tmp_path / "missing" / "x.xlsx")), "No such folder"),
                      (("init", str(tmp_path / "none.xlsx")), "No such tracker")):
        out = run(*args)
        assert out.returncode == 2 and msg in out.stderr and "Traceback" not in out.stderr

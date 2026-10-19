"""Getting started: write a config for an existing tracker, or create a new one.

    tracker-mcp init my-tracker.xlsx            # guess the config from headers and dropdowns
    tracker-mcp init my-tracker.csv             # a CSV export (Notion, Huntr, Teal...)
    tracker-mcp init https://docs.google.com/spreadsheets/d/<id>/edit
    tracker-mcp new my-tracker.xlsx             # an empty tracker plus its config (.xlsx or .csv)

`init` reads the tracker and guesses: which sheet and row hold the headers, which
header is which field, what each status means, and how text dates are written.
It writes the guess as a commented TOML file and marks with `# CHECK` whatever it
was unsure about. Nothing in the tracker changes. The patterns understand English
and Russian; anything else ends up in the "check by hand" list, never in a
silent default.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .config import ConfigError, config_path, load_config
from .stores import is_google_sheet, open_store
from .tracker import Tracker, TrackerError, dropdown_values, parse_date, real_title, sheet_in

# Text date formats init tries when a date column holds text the built-in parser
# does not read (ISO, dd.mm.yyyy and spreadsheet dates are always read).
DATE_CANDIDATES = ["%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y", "%Y/%m/%d", "%b %d, %Y", "%B %d, %Y",
                   "%d %b %Y", "%d %B %Y", "%m-%d-%Y", "%d-%m-%Y", "%b %d %Y", "%d.%m.%y"]
# Statuses added when the tracker has no dropdown and its data lacks a kind, so the
# server can still set them. By language of the headers.
STANDARD = {
    "en": {"not_applied": "Not applied", "applied": "Applied", "screening": "Screening",
           "interview": "Interview", "offer": "Offer", "rejected": "Rejected", "withdrawn": "Withdrawn",
           "no_response": "No response"},
    "ru": {"not_applied": "Не подавался", "applied": "Подана", "screening": "Скрининг",
           "interview": "Интервью", "offer": "Оффер", "rejected": "Отказ", "withdrawn": "Отозвал",
           "no_response": "Ответа нет"},
}

# Field -> header pattern (case-insensitive). Order matters: a header goes to the
# first field that claims it, so the specific follow-up fields come before
# "status" and "contact", and the date of applying before the yes/no column.
FIELD_PATTERNS: list[tuple[str, str]] = [
    ("num", r"^(#|№|no\.?|nr\.?|n°|id|п/п|номер)$"),
    ("reached", r"furthest|highest stage|max(imum)? stage|stage reached|^reached|дош[её]л|дошли|макс\w* этап"),
    ("followup_on", r"(follow.?up|nudge|reminder|фоллоу.?ап|напомина).*(date|\bon\b|when|due|next|дата|когда)"
                    r"|^(next )?follow.?up( 1)?$|^фоллоу.?ап( 1)?$"),
    ("followup_contact", r"(follow.?up|фоллоу.?ап).*(contact|who|\bto\b|person|кому|контакт)"
                         r"|^contact|recruiter|hiring manager|контакт|^кому"),
    ("followup_log", r"(follow.?up|фоллоу.?ап).*(log|status|history|notes|статус|лог|истори)"
                     r"|^(log|history|история|лог)$"),
    ("applied_on", r"(date|when|дата).*(appl|submit|sent|подач|отклик|отправ)"
                   r"|(appl|submit|sent|подан|отклик).*(date|\bon\b|when|дата)"),
    ("applied", r"^(applied|submitted|sent|подался|подалась|подано|откликнулся|отправлено)\s*\??$"),
    ("company", r"company|employer|organi[sz]ation|компани|работодател|фирма|организац"),
    ("role", r"^(job )?(title|role|position)\b|^job$|^(vacancy|opening)\b|роль|должност|позици|ваканси"),
    ("status", r"status|^stage$|^state$|статус|^этап$|^стадия$"),
    ("link", r"\blink\b|\burl\b|posting|ссылк"),          # not "LinkedIn"
    ("channel", r"channel|applied (via|through)|^via$|канал|как подал"),
    ("source", r"source|found (on|via|at)|where found|источник|где наш[её]л|откуда"),
    ("response", r"^(response|reply|replies|answer|feedback)\b(?!.*(date|when))|^ответ(?!.*дат)"),
    ("notes", r"^notes?\b|comment|заметк|комментар|примечан"),
]

# What a status word means. Checked in this order; the first match wins. "final"
# is a status after the last interview ("Waiting for offer"): an interview stage,
# not a waiting status, so close_stale never touches it.
STATUS_PATTERNS: list[tuple[str, str]] = [
    ("withdrawn", r"withdr|not interested|no longer interested|lost interest|отозвал|отказался|отказалась"
                  r"|не интересн"),
    ("not_applied", r"^(not (yet )?applied|to apply|wish ?list|saved|bookmarked?|draft|interested|planned"
                    r"|planning to apply|researching|backlog)$|не подав|не подан|^хочу|^план"),
    ("no_response", r"no (response|reply|answer)|ghost|нет ответа|ответа нет|без ответа|молч"),
    ("final", r"(wait|await|pending).*(offer|decision)|(offer|decision) pending|жд[её]м (оффер|решени)"),
    ("offer", r"offer|hired|accepted|оффер|принят"),        # "Offer declined" got an offer too
    ("rejected", r"reject|declin|turned down|not selected|unsuccessful|отказ"),
    ("waiting", r"wait|pending|awaiting|under review|in review|жд[её]м|ожида|на рассмотрени"),
    ("applied", r"^(applied|submitted|sent|application (sent|submitted))$|^подан[аоы]?$|^отправлен|^откликнул"),
    ("closed", r"closed|filled|cancel|archiv|закрыт|отмен"),
    ("screening", r"screen|recruiter|\bhr\b|phone|intro|скрининг|созвон|hr-"),
    ("interview", r"interview|tech|onsite|on-site|final|panel|case|assignment|take.?home|test|assessment"
                  r"|challenge|интервью|тех|финал|задани|собес"),
]
STAGE_KINDS = {"screening": ["screening"], "interview": ["interview", "final"], "offer": ["offer"]}

YES = r"^(yes|y|да|true|1|✓|✔|x)$"

MESSAGES = {
    "en": {"applied": "applied", "applied_with_followup": "applied, follow-up 1 on {date}",
           "followup_cancelled": "follow-up {date} cancelled: {status}",
           "stale_closed": "closed: no reply for {days} days"},
    "ru": {"applied": "подано", "applied_with_followup": "подано, фоллоу-ап 1 — {date}",
           "followup_cancelled": "фоллоу-ап {date} снят — {status}",
           "stale_closed": "закрыто: {days} дн. без ответа"},
}


@dataclass
class Guess:
    sheet: str
    header_row: int
    rows: int
    columns: dict[str, tuple[int, str]]          # field -> (column, header)
    unmapped: list[tuple[int, str]]
    statuses: list[str]
    statuses_from: str                           # "dropdown", "data" or "none"
    meaning: dict[str, list[str]]                # not_applied, applied, waiting, ... -> statuses
    yes_no: tuple[str, str] | None
    reached_words: dict[str, list[str]]          # stage -> words of the reached column's own dropdown
    summary_sheet: str | None
    lang: str
    date_formats: list[str] = field(default_factory=list)
    checks: list[str] = field(default_factory=list)


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def _match_fields(headers: dict[int, str]) -> dict[str, tuple[int, str]]:
    taken: set[int] = set()
    out: dict[str, tuple[int, str]] = {}
    for name, pattern in FIELD_PATTERNS:
        rx = re.compile(pattern)
        for col, text in headers.items():
            if col not in taken and rx.search(_norm(text)):
                out[name] = (col, text)
                taken.add(col)
                break
    return out


def _headers(ws: Worksheet, row: int) -> dict[int, str]:
    return {c: str(v).strip() for c in range(1, ws.max_column + 1)
            if isinstance(v := ws.cell(row, c).value, str) and v.strip()}


def _find_header(wb, sheet: str | None) -> tuple[Worksheet, int, dict[str, tuple[int, str]]]:
    if sheet:
        sheet = sheet_in(wb, sheet)
    if sheet and sheet not in wb.sheetnames:
        raise TrackerError(f"No sheet «{sheet}». Sheets: {', '.join(real_title(w) for w in wb.worksheets)}")
    best = None
    for ws in [wb[sheet]] if sheet else wb.worksheets:
        for r in range(1, min(ws.max_row, 30) + 1):
            fields = _match_fields(_headers(ws, r))
            if "company" not in fields or not ({"role", "status"} & fields.keys()):
                continue
            if best is None or len(fields) > len(best[2]):
                best = (ws, r, fields)
    if best is None:
        where = f"sheet «{sheet}»" if sheet else "any sheet"
        raise TrackerError(f"No header row found in {where}: looked for a company column next to a "
                           "role or status column in the first 30 rows. Write the config by hand: "
                           "`tracker-mcp example` prints a commented one.")
    return best


def _kind(status: str, kinds: list[tuple[str, str]] = STATUS_PATTERNS) -> str | None:
    return next((name for name, p in kinds if re.search(p, _norm(status))), None)


def _classify(statuses: list[str]) -> tuple[dict[str, list[str]], list[str]]:
    meaning: dict[str, list[str]] = {name: [] for name, _ in STATUS_PATTERNS}
    stages_only = [(k, p) for k, p in STATUS_PATTERNS if k in ("screening", "interview")]
    unknown = []
    for s in statuses:
        kind = _kind(s)
        if kind is None:
            unknown.append(s)
            continue
        meaning[kind].append(s)
        if kind == "rejected" and (stage := _kind(s, stages_only)):  # "Rejected after interview"
            meaning[stage].append(s)
    return meaning, unknown


def _one_line(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _date_formats(samples: list[str], lang: str) -> tuple[list[str], list[str], list[str]]:
    """Formats that read the text dates the built-in parser does not, the values no
    format reads, and notes on ambiguity (03/04/2026: March or April?)."""
    todo = [s for s in samples if parse_date(s) is None]
    if not todo:
        return [], [], []
    scores = []
    for fmt in DATE_CANDIDATES:
        ok = sum(parse_date(s, [fmt]) is not None for s in todo)
        if ok:
            scores.append((ok, fmt))
    if not scores:
        return [], todo, []
    best = max(ok for ok, _ in scores)
    picks = [fmt for ok, fmt in scores if ok == best]
    notes = []
    if "%m/%d/%Y" in picks and "%d/%m/%Y" in picks:   # every date has both parts <= 12
        pick = "%d/%m/%Y" if lang == "ru" else "%m/%d/%Y"
        notes.append(f"dates like «{todo[0]}» could be month-first or day-first; assumed "
                     f"{'day' if pick.startswith('%d') else 'month'}-first ({pick})")
    else:
        pick = picks[0]
    failed = [s for s in todo if parse_date(s, [pick]) is None]
    return [pick], failed, notes


def guess(path: str | Path, sheet: str | None = None) -> Guess:
    store = open_store(path)
    wb, _, _ = store.load()
    if store.single_sheet:
        sheet = None
    ws, header_row, fields = _find_header(wb, sheet)
    headers = _headers(ws, header_row)
    company_col = fields["company"][0]
    data_rows = [r for r in range(header_row + 1, ws.max_row + 1)
                 if ws.cell(r, company_col).value not in (None, "")]
    checks: list[str] = []

    def values(col: int) -> list:
        return [v for r in data_rows if (v := ws.cell(r, col).value) not in (None, "")]

    # "Applied" holding dates is the date of applying, not a yes/no flag.
    if "applied" in fields and "applied_on" not in fields:
        vals = values(fields["applied"][0])
        if vals and sum(parse_date(v) is not None for v in vals) > len(vals) / 2:
            fields["applied_on"] = fields.pop("applied")
    yes_no = None
    if "applied" in fields:
        raw = values(fields["applied"][0])
        options = dropdown_values(ws, fields["applied"][0]) or list(dict.fromkeys(str(v).strip() for v in raw))
        yes = [o for o in options if re.match(YES, _norm(o))]
        no = [o for o in options if o not in yes]
        if any(isinstance(v, bool) for v in raw):   # checkboxes: the server reads and writes booleans
            yes_no = ("TRUE", "FALSE")
        elif len(yes) == 1 and len(no) == 1:
            yes_no = (yes[0], no[0])
        else:
            checks.append(f"values.yes / values.no: could not tell yes from no in "
                          f"«{fields['applied'][1]}» ({', '.join(options) or 'empty'})")

    status_col = fields["status"][0] if "status" in fields else None
    if status_col is None:
        raise TrackerError("No status column found. Write the config by hand: "
                           "`tracker-mcp example` prints a commented one.")
    statuses, statuses_from = dropdown_values(ws, status_col), "dropdown"
    if not statuses:
        statuses = list(dict.fromkeys(str(v).strip() for v in values(status_col)))
        statuses_from = "data" if statuses else "none"
    meaning, unknown = _classify(statuses)
    if statuses_from == "none":
        checks.append("the status column has no dropdown and no values yet: the default English "
                      "statuses apply (see `fallback` in the example config)")
    reached_words: dict[str, list[str]] = {}
    if "reached" in fields:   # its own words, e.g. "Final round", join the stages
        col = fields["reached"][0]
        words = dropdown_values(ws, col) or list(dict.fromkeys(str(v).strip() for v in values(col)))
        for w in words:
            kind = _kind(w)
            stage = next((st for st, kinds in STAGE_KINDS.items() if kind in kinds), None)
            if stage:
                reached_words.setdefault(stage, []).append(w)
            else:
                checks.append(f"«{w}» in «{fields['reached'][1]}» is not a screening, interview or offer "
                              "word; add it to a stage in [statuses.stages]")
    if unknown:
        checks.append(f"statuses not recognised: {', '.join(unknown)}. Add them to closed, waiting "
                      "or a funnel stage if they mean one.")
    for kind, what in (("not_applied", "not_applied (an application not sent yet)"),
                       ("applied", "applied (the status right after applying)"),
                       ("no_response", "no_response (what close_stale sets)")):
        if not meaning[kind]:
            checks.append(f"statuses.{what}: no status looks like it")
        elif len(meaning[kind]) > 1:
            checks.append(f"statuses.{kind}: several candidates ({', '.join(meaning[kind])}), took the first")

    title = real_title(ws)
    summary = next((real_title(other) for other in wb.worksheets if other is not ws and any(
        isinstance(c.value, str) and c.value.startswith("=") and c.data_type == "f"
        and (f"{title}!" in c.value or f"'{title}'!" in c.value)
        for row in other.iter_rows() for c in row)), None)
    lang = "ru" if any(re.search("[а-яё]", h, re.I) for h in headers.values()) else "en"
    mapped = {c for c, _ in fields.values()}

    # Text dates: find the format, report what no format reads.
    date_formats: list[str] = []
    texts = [str(v).strip() for f in ("applied_on", "followup_on") if f in fields
             for v in values(fields[f][0]) if isinstance(v, str)]
    if texts:
        date_formats, failed, notes = _date_formats(texts, lang)
        checks.extend(notes)
        if failed:
            sample = ", ".join(f"«{x}»" for x in failed[:3])
            checks.append(f"{len(failed)} value(s) in the date columns are not dates the server can read "
                          f"({sample}); fix them in the tracker or add a format to values.date_formats")

    # Columns the follow-up tools need.
    if "followup_on" not in fields:
        hint = next((h for c, h in headers.items()
                     if c not in mapped and re.search(r"next|step|follow|фоллоу", h, re.I)), None)
        checks.append("no follow-up date column: followups_due, log_followup and close_stale are off. Add a "
                      "column such as «Follow-up date»" + (f" («{hint}» holds text, not dates)" if hint else "")
                      + " and run init again")
    elif "followup_log" not in fields:
        checks.append("no follow-up log column: log_followup is off. Add a column such as «Follow-up log»")
    if "followup_contact" in fields and not re.search(r"follow|фоллоу", fields["followup_contact"][1], re.I):
        checks.append(f"«{fields['followup_contact'][1]}» is used as the follow-up contact: "
                      "log_followup(contact=...) replaces what is there")

    # No dropdown: add the standard statuses the data does not have yet.
    if statuses_from != "dropdown":
        added = []
        for kind, word in STANDARD[lang].items():
            if not meaning[kind] and word not in statuses:
                meaning[kind].append(word)
                statuses.append(word)
                added.append(word)
        if added:
            checks.append(f"no status dropdown: added {', '.join(added)} to the allowed statuses (fallback)")
        checks[:] = [c for c in checks if not c.startswith(("statuses.not_applied", "statuses.applied (",
                                                             "statuses.no_response"))]
    return Guess(
        sheet=real_title(ws), header_row=header_row, rows=len(data_rows),
        columns=dict(sorted(fields.items(), key=lambda kv: kv[1][0])),
        unmapped=[(c, h) for c, h in headers.items() if c not in mapped],
        statuses=statuses, statuses_from=statuses_from, meaning=meaning,
        yes_no=yes_no, reached_words=reached_words, summary_sheet=summary, lang=lang,
        date_formats=date_formats, checks=[_one_line(c) for c in checks])


# ---------------------------------------------------------------- TOML

def _shell(text: str) -> str:
    return f'"{text}"' if re.search(r"[\s&?#]", text) else text


def _q(s: str) -> str:
    return json.dumps(s, ensure_ascii=False)


def _list(items: list[str]) -> str:
    return "[" + ", ".join(_q(s) for s in items) + "]"


def _key(text: str, col: int) -> str:
    if re.search(r"[^\x00-\x7f]", text):   # a non-Latin header gets the column letter
        return f"col_{get_column_letter(col).lower()}"
    slug = re.sub(r"[^a-z0-9]+", "_", text.casefold()).strip("_")
    return slug if re.fullmatch(r"[a-z][a-z0-9_]{0,30}", slug or "") else f"col_{get_column_letter(col).lower()}"


def render(g: Guess, target: str = "<tracker>") -> str:
    m = g.meaning
    first = lambda kind, default: m[kind][0] if m[kind] else default  # noqa: E731
    not_applied = first("not_applied", "Not applied")
    applied = first("applied", "Applied")
    no_response = first("no_response", "No response")
    rejected = m["rejected"]
    closed = [s for s in g.statuses if s in {*m["withdrawn"], *m["no_response"], *m["rejected"],
                                             *m["offer"], *m["closed"]}]
    waiting = [applied, *m["waiting"]] if m["applied"] else m["waiting"] or [applied]
    stages = {}
    for stage, kinds in STAGE_KINDS.items():
        words = [s for kind in kinds for s in m[kind]] + g.reached_words.get(stage, [])
        if words:
            stages[stage] = list(dict.fromkeys(words))
    check = " # CHECK" if g.checks else ""
    yes, no = g.yes_no or ("Yes", "No")
    lines = [
        "# tracker-mcp config, written by `tracker-mcp init`. Lines marked CHECK are guesses;",
        "# `tracker-mcp example` prints a config with every key explained.",
        f"# See how it maps the tracker:  tracker-mcp --check --tracker {_shell(target)}",
        "",
        f"sheet = {_q(g.sheet)}",
        f"header_row = {g.header_row}",
    ]
    if g.summary_sheet:
        lines.append(f"summary_sheet = {_q(g.summary_sheet)}   # its formulas grow with new rows")
    lines.append("date_columns = " + _list([f for f in ("applied_on", "followup_on") if f in g.columns]))
    lines += ["", "[columns]                       # field = header text (exact)"]
    for name, (col, header) in g.columns.items():
        line = f"{name} = {_q(header)}"
        lines.append(f"{line.ljust(39)} # {get_column_letter(col)}")
    if g.unmapped:
        lines += ["# Not mapped; returned under \"other\". Uncomment to give one a field name:"]
        lines += [f"# {_key(h, c)} = {_q(h)}" for c, h in g.unmapped]
    lines += ["", f"[statuses]                      # statuses from the {g.statuses_from}"]
    lines += [
        f"not_applied = {_q(not_applied)}" + ("" if m["not_applied"] else "   # CHECK"),
        f"applied = {_q(applied)}" + ("" if m["applied"] else "   # CHECK"),
        f"waiting = {_list(waiting)}",
        f"no_response = {_q(no_response)}" + ("" if m["no_response"] else "   # CHECK"),
        f"rejected = {_list(rejected)}",
        f"closed = {_list(closed)}",
    ]
    if g.statuses_from != "dropdown":
        lines.append(f"fallback = {_list(g.statuses)}   # the status column has no dropdown")
    lines += ["", "[statuses.stages]               # funnel stages after a reply, in order"]
    lines += [f"{k} = {_list(v)}" for k, v in stages.items()]
    lines += ["", "[values]",
              f"yes = {_q(yes)}" + ("" if g.yes_no or "applied" not in g.columns else "   # CHECK"),
              f"no = {_q(no)}" + ("" if g.yes_no or "applied" not in g.columns else "   # CHECK"),
              'no_date = "—"',
              'separator = " | "',
              f"stamp_format = {_q('%d.%m' if g.lang == 'ru' else '%Y-%m-%d')}",
              *([f"date_formats = {_list(g.date_formats)}   # how this tracker writes text dates"]
                if g.date_formats else []),
              "", "[followups]", "first_after_days = 10", "next_after_days = 14", "weekend_to_monday = true",
              "", "[messages]                      # appended to the follow-up log"]
    lines += [f"{k} = {_q(v)}" for k, v in MESSAGES[g.lang].items()]
    lines += ["", "[safety]", "confirm_writes = true", ""]
    if check:
        lines.insert(3, "# CHECK by hand:\n" + "\n".join(f"#   - {c}" for c in g.checks))
    return "\n".join(lines)


def report(g: Guess, out: Path | None) -> str:
    lines = [f"Wrote {out}" if out else "Config (not written):",
             f"Sheet «{g.sheet}», header row {g.header_row}, {g.rows} application rows."]
    if g.summary_sheet:
        lines.append(f"Summary sheet: «{g.summary_sheet}».")
    lines.append("Columns:")
    lines += [f"  {name:<17}{get_column_letter(col):<4}{h}" for name, (col, h) in g.columns.items()]
    if g.unmapped:
        lines.append("Not mapped (returned under \"other\"): " + ", ".join(h for _, h in g.unmapped))
    meaning = {k: v for k, v in g.meaning.items() if v}
    lines.append(f"Statuses ({g.statuses_from}): " + "; ".join(f"{k}: {', '.join(v)}" for k, v in meaning.items()))
    if g.checks:
        lines.append("Check by hand (marked CHECK in the file):")
        lines += [f"  - {c}" for c in g.checks]
    return "\n".join(lines)


def _target(path: str | Path) -> str:
    t = str(path)
    return t if is_google_sheet(t) else str(Path(t).expanduser().resolve())


def init(path: str | Path, *, out: str | Path | None = None, sheet: str | None = None,
         force: bool = False, write: bool = True) -> tuple[Path | None, Guess, str]:
    t = _target(path)
    if not is_google_sheet(t) and not Path(t).exists():
        raise TrackerError(f"No such tracker: {t}")
    g = guess(t, sheet)
    text = render(g, t)
    target = Path(out).expanduser() if out else config_path(t)
    if write and target.exists() and not force:
        raise TrackerError(f"{target} already exists. Pass --force to overwrite it, "
                           "or --print to only see the guess.")
    try:
        check(t, text)
    except TrackerError as e:
        if write:
            raise
        text = f"# DOES NOT WORK YET: {_one_line(str(e))}\n{text}"
    if write:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")
    return (target if write else None), g, text


def check(target: str, text: str) -> None:
    """The guess must load and map the tracker before it is written anywhere."""
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d) / "guess.tracker-mcp.toml"
        tmp.write_text(text, encoding="utf-8")
        try:
            Tracker(target, config=load_config(target, tmp), backup_dir=d).describe()
        except (ConfigError, TrackerError) as e:
            raise TrackerError(f"The guessed config does not work: {e}. Nothing was written; "
                               "see the guess with --print and fix it by hand.") from e


# ---------------------------------------------------------------- new

TEMPLATE_HEADERS = ["#", "Company", "Role", "Link", "Source", "Applied?", "Applied on", "Channel", "Status",
                    "Furthest stage", "Response", "Follow-up date", "Follow-up contact", "Follow-up log",
                    "Notes"]
TEMPLATE_STATUSES = ["Not applied", "Applied", "Screening", "Interview", "Offer", "Rejected", "Withdrawn",
                     "No response", "Closed"]
TEMPLATE_WIDTHS = [5, 22, 32, 30, 16, 10, 12, 18, 14, 14, 30, 14, 24, 40, 40]


def new_tracker(path: str | Path, *, force_config: bool = False) -> tuple[Path, Path]:
    """An empty tracker that every tool works with, plus its config (made by init).
    .xlsx gets dropdowns, a filter and a summary sheet; .csv gets the header row."""
    path = Path(path).expanduser().resolve()
    if path.suffix.lower() not in (".xlsx", ".csv"):
        raise TrackerError("The tracker must be an .xlsx or a .csv file. For Google Sheets, import the "
                           ".xlsx that `new` makes (File → Import), then run `tracker-mcp init <sheet URL>`.")
    if not path.parent.is_dir():
        raise TrackerError(f"No such folder: {path.parent}")
    if path.exists():
        raise TrackerError(f"{path} already exists; `new` never overwrites a tracker. "
                           "For an existing tracker use `tracker-mcp init`.")
    if config_path(path).exists() and not force_config:
        raise TrackerError(f"{config_path(path)} already exists. Pass --force to overwrite it.")
    if path.suffix.lower() == ".csv":
        path.write_text(",".join(TEMPLATE_HEADERS) + "\n", encoding="utf-8")
        target, _, _ = init(path, force=True)
        return path, target
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(TEMPLATE_HEADERS)
    for c, width in enumerate(TEMPLATE_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = width
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8E8E8")
    ws.freeze_panes = "A2"
    last = get_column_letter(len(TEMPLATE_HEADERS))
    ws.auto_filter.ref = f"A1:{last}1"
    col = {h: get_column_letter(i) for i, h in enumerate(TEMPLATE_HEADERS, 1)}
    for header, options in (("Applied?", ["Yes", "No"]), ("Status", TEMPLATE_STATUSES),
                            ("Furthest stage", ["Screening", "Interview", "Offer"])):
        dv = DataValidation(type="list", formula1=_q(",".join(options)), allow_blank=True)
        dv.add(f"{col[header]}2:{col[header]}500")
        ws.add_data_validation(dv)
    s = wb.create_sheet("Summary")
    rows = [("Applications", f"=COUNTA(Applications!{col['Company']}2:{col['Company']}500)"),
            ("Submitted", f'=COUNTIF(Applications!{col["Applied?"]}2:{col["Applied?"]}500,"Yes")'),
            ("Waiting", f'=COUNTIF(Applications!{col["Status"]}2:{col["Status"]}500,"Applied")'),
            ("Reached an interview", f'=COUNTIF(Applications!{col["Furthest stage"]}2:'
                                     f'{col["Furthest stage"]}500,"Interview")'
                                     f'+COUNTIF(Applications!{col["Furthest stage"]}2:'
                                     f'{col["Furthest stage"]}500,"Offer")')]
    for r, (label, formula) in enumerate(rows, 1):
        s.cell(r, 1, label).font = Font(bold=True)
        s.cell(r, 2, formula)
    s.column_dimensions["A"].width = 22
    wb.save(path)
    target, _, _ = init(path, force=True)
    return path, target


def example_config() -> str:
    """The commented example config: packaged with the wheel, or from a clone."""
    from importlib import resources
    try:
        return (resources.files("tracker_mcp") / "example.tracker-mcp.toml").read_text(encoding="utf-8")
    except (FileNotFoundError, OSError):
        repo = Path(__file__).resolve().parents[2] / "examples" / "tracker-mcp.example.toml"
        return repo.read_text(encoding="utf-8")


# ---------------------------------------------------------------- CLI

def _next(target: str) -> str:
    return (f"\nNext, check the mapping:  tracker-mcp --check --tracker {_shell(target)}\n"
            "(from a clone of the repository: uv run tracker-mcp ...)")


def cli(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(prog="tracker-mcp")
    sub = parser.add_subparsers(dest="cmd", required=True)
    p_init = sub.add_parser("init", help="write a config for an existing tracker (.xlsx, .csv, Google Sheet URL)")
    p_init.add_argument("tracker")
    p_init.add_argument("--sheet", help="the sheet with one row per application (default: guessed)")
    p_init.add_argument("--config", help="where to write (default: <tracker>.tracker-mcp.toml next to it; "
                                         "for a Google Sheet ~/.tracker-mcp/<id>.tracker-mcp.toml)")
    p_init.add_argument("--force", action="store_true", help="overwrite an existing config")
    p_init.add_argument("--print", dest="print_only", action="store_true",
                        help="print the config instead of writing it")
    p_new = sub.add_parser("new", help="create an empty tracker (.xlsx or .csv) and its config")
    p_new.add_argument("tracker")
    p_new.add_argument("--force", action="store_true", help="overwrite an existing config")
    sub.add_parser("example", help="print a config with every key explained")
    p_auth = sub.add_parser("auth", help="sign in to Google once, for trackers in Google Sheets")
    p_auth.add_argument("provider", choices=["google"])
    p_auth.add_argument("--client-secret", required=True,
                        help="client_secret.json of an OAuth client of type Desktop app (Google Cloud console)")
    p_auth.add_argument("--no-browser", action="store_true", help="only print the sign-in link")
    args = parser.parse_args(argv)
    try:
        if args.cmd == "init":
            target, g, text = init(args.tracker, out=args.config, sheet=args.sheet, force=args.force,
                                   write=not args.print_only)
            print(text if args.print_only else report(g, target) + _next(_target(args.tracker)))
        elif args.cmd == "new":
            book, cfg = new_tracker(args.tracker, force_config=args.force)
            print(f"Created {book}\nand its config {cfg}" + _next(str(book)))
        elif args.cmd == "example":
            print(example_config())
        else:
            from .gsheets import login
            token = login(args.client_secret, open_browser=not args.no_browser)
            print(f"Signed in. The token is in {token} (keep it private).\n"
                  "Next: tracker-mcp init https://docs.google.com/spreadsheets/d/<id>/edit")
    except TrackerError as e:
        print(f"tracker-mcp: {e}", file=sys.stderr)
        return 2
    return 0

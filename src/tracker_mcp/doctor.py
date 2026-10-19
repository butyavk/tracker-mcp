"""check_tracker and fix_tracker: what hand edits leave behind, found and fixed in one write.

A tracker kept by hand, or by a model writing a script for each change, drifts: dates
saved as text in one session and as dates in another, rows added without a status, a
follow-up on a Sunday, a rejection with no reply date. Each check here is a rule the
write tools already follow; this module finds the rows that break one. Every issue
says what is wrong, and most carry the value that fixes it. The rest need a decision
(an unreadable date, a duplicate) and are only reported.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import TYPE_CHECKING, Any

from openpyxl.cell.cell import MergedCell

from .tracker import _ANY_DATE, Layout, TrackerError, company_key, date_in_text, undid, weekday_or_monday

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

    from .tracker import Tracker

NOFIX = object()

KINDS = {
    "date_as_text": "A date kept as text or as a bare number: Excel cannot sort or filter it as a date",
    "bad_date": "Not a readable date in a date column",
    "date_with_text": "A date followed by other text in a date column: split it by hand, so nothing is lost",
    "no_status": "No status: the row drops out of status filters and counts",
    "unknown_status": "A status that is not in the dropdown (or spelled differently)",
    "weekend_followup": "A follow-up date on a weekend; the follow-up rule moves them to Monday",
    "closed_with_followup": "A closed application with a follow-up still planned",
    "no_followup": "Submitted and waiting, with no follow-up date and no decision to skip it",
    "applied_mismatch": "The applied column disagrees with the status",
    "reached_behind": "The furthest-stage column is behind the current status",
    "no_reply_date": "A reply with no reply date",
    "duplicate": "The same company and role on two rows",
}

# Clauses of a response cell about the application itself (submitted, received, an
# auto-reply), not a reply from the company.
_NOT_A_REPLY = re.compile(
    r"автоподтвержд|автоответ|автописьм|экранное подтвержд|подтвержден\w*\s+(от|заявк|подач|получен|по почте)"
    r"|^\W*подтвержд|auto.?(reply|confirm|response)|acknowledg|confirm\w*\s+(of |that )?(your |the )?"
    r"(application|receipt|submission)|application confirm|received your|we.ve received|we have received"
    r"|thank(s| you) for (applying|your (application|interest))|application (was )?received|получили"
    r"|\b(applied|submitted|sent (my|the|an?) (cv|application|resume))\b|подал(ся|ась|и)?\b|подан[аоы]?\b"
    r"|подача|отправ\w* (cv|резюме|заявк)",
    re.IGNORECASE)
# A date cell that is only a date: ISO (with an optional time), dd.mm.yyyy, or a serial.
_WHOLE_DATE = re.compile(r"\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?|\d{1,2}\.\d{1,2}\.\d{4}|\d{5}")


@dataclass
class Issue:
    kind: str
    row: int
    field: str | None
    value: Any
    note: str
    fix: Any = NOFIX

    @property
    def fixable(self) -> bool:
        return self.fix is not NOFIX


def reply_dates_from_audit(t: Tracker, labels: dict[int, str]) -> dict[int, date]:
    """Rows whose first reply status was set through this server, and the day it was set.
    An entry counts only if it names the row as it is now (`labels`: row -> "Company —
    Role"; rows move when the sheet is sorted) and was not undone later."""
    entries = t._audit_entries()
    undos = [e for e in entries if e.get("tool") == "undo_last_write"]
    out: dict[int, date] = {}
    for e in entries:
        if e.get("tool") != "set_status" or e.get("unverified") or any(undid(u, e) for u in undos):
            continue
        for ch in e.get("changes", []):
            row = int(ch.get("row", 0))
            if (ch.get("field") == "status" and t._is_reply(ch.get("new")) and not t._is_reply(ch.get("old"))
                    and (e.get("labels") or {}).get(str(row), (e.get("labels") or {}).get(row)) == labels.get(row)
                    and row in labels):
                try:
                    out.setdefault(row, datetime.fromisoformat(e["at"]).date())
                except (KeyError, ValueError):
                    continue
    return out


def _parses(text: str, fmt: str) -> bool:
    try:
        datetime.strptime(text, fmt)
        return True
    except ValueError:
        return False


def reply_date_hint(t: Tracker, rec: dict[str, Any], audit_day: date | None) -> tuple[date | None, str | None]:
    """The day the first reply came, as far as the row tells: the earliest date in the
    response column outside the clauses about the application being received, else the
    day this server recorded the reply."""
    applied = t._date(rec.get("applied_on"))
    today = t.today()
    text = str(rec.get("response") or "")
    found = []
    for clause in re.split(r";|\||\.\s|\n", text):
        if _NOT_A_REPLY.search(clause):
            continue
        for m in _ANY_DATE.finditer(clause):
            d = date_in_text(m.group(1), today, applied)
            if d and (applied is None or d > applied):   # a reply comes after the submit day
                found.append((d, clause.strip()))
    if found:
        d, evidence = min(found, key=lambda x: x[0])
        return d, f"from the response column: «{evidence[:90]}»"
    if audit_day:
        return audit_day, "the day this server recorded the reply"
    return None, None


def find_issues(t: Tracker, ws: Worksheet, lay: Layout) -> list[Issue]:
    cfg, st = t.cfg, t.cfg.statuses
    today = t.today()
    allowed = t._allowed_statuses(ws, lay)
    by_fold = {s.casefold(): s for s in allowed}
    audit_days = (reply_dates_from_audit(t, {r: t._label(ws, lay, r) for r in lay.rows()})
                  if cfg.has("replied_on") else {})
    seen: dict[tuple[str, str], int] = {}
    out: list[Issue] = []

    def word(status: str) -> str:   # the dropdown's own spelling of a status
        return by_fold.get(status.casefold(), status)

    for r in lay.rows():
        if ws.cell(r, lay.cols["company"]).value in (None, ""):
            continue
        if any(isinstance(ws.cell(r, c), MergedCell) for c in lay.cols.values()):
            continue   # a merged separator or heading row, not an application
        rec = t._record(ws, lay, r)
        raw_status = ws.cell(r, lay.cols["status"]).value
        status = str(raw_status).strip() if raw_status not in (None, "") else None

        # dates: a CSV has no cell types, so there everything is text and nothing to fix
        if t.store.kind != "csv":
            for f in (f for f in lay.cols if cfg.is_date(f)):
                v = ws.cell(r, lay.cols[f]).value
                if v in (None, "") or isinstance(v, (datetime, date)) or str(v).strip() == cfg.values.no_date:
                    continue
                d = t._date(v)
                whole = isinstance(v, (int, float)) or _WHOLE_DATE.fullmatch(str(v).strip()) or any(
                    _parses(str(v).strip(), fmt) for fmt in cfg.values.date_formats)
                if d and whole:
                    out.append(Issue("date_as_text", r, f, v, f"«{v}» is kept as text or a number, not a date",
                                     fix=d))
                elif d:
                    out.append(Issue("date_with_text", r, f, v, f"«{v}»: a date with more text after it; move "
                                                                "the text to the notes and leave the date"))
                else:
                    out.append(Issue("bad_date", r, f, v, f"«{v}» is not a date the server can read; "
                                                          "fix it by hand or add its format to values.date_formats"))

        # status
        if status is None:
            flag = rec.get("applied")
            flag_text = str(flag).strip().casefold() if flag not in (None, "") else ""
            has_date = t._date(rec.get("applied_on")) is not None
            if flag is True or flag_text == cfg.values.yes.casefold():
                applied, why = True, "the applied column"
            elif flag is False or flag_text == cfg.values.no.casefold():
                applied, why = False, "the applied column"
            else:
                applied, why = has_date, "the submit date" if has_date else "no submit date"
            new = word(st.applied if applied else st.not_applied)
            out.append(Issue("no_status", r, "status", None, f"no status; «{new}» from {why}", fix=new))
        elif status.casefold() not in by_fold:
            alias = t.aliases.get(status.casefold())
            fix = alias if alias in allowed else NOFIX
            out.append(Issue("unknown_status", r, "status", status,
                             f"«{status}» is not in the dropdown" + (f"; it means «{alias}»" if fix is not NOFIX
                                                                     else "; pick one with set_status"), fix=fix))
            if fix is not NOFIX:
                status = alias   # the later checks see what the status means
        elif status != by_fold[status.casefold()]:
            out.append(Issue("unknown_status", r, "status", status,
                             f"spelled differently from the dropdown («{by_fold[status.casefold()]}»)",
                             fix=by_fold[status.casefold()]))
        status = word(status) if status else None

        # follow-ups
        if cfg.has("followup_on"):
            fu_raw = ws.cell(r, lay.cols["followup_on"]).value
            fu = t._date(fu_raw)
            is_open = status not in t.closed and status != st.not_applied
            if fu and fu >= today and status in t.closed:
                out.append(Issue("closed_with_followup", r, "followup_on", fu,
                                 f"closed ({status}) but a follow-up is planned for {fu.isoformat()}",
                                 fix=cfg.values.no_date))
            elif fu and fu >= today and is_open and cfg.followups.weekend_to_monday and fu.weekday() >= 5:
                monday = weekday_or_monday(fu)
                out.append(Issue("weekend_followup", r, "followup_on", fu,
                                 f"{fu.isoformat()} is a {fu.strftime('%A')}; moved to {monday.isoformat()}",
                                 fix=monday))
            if (is_open and status in t.waiting and t._is_applied(rec) and fu_raw in (None, "")):
                out.append(Issue("no_followup", r, "followup_on", None,
                                 "waiting with no follow-up date: plan one (reschedule_followup) or write "
                                 f"«{cfg.values.no_date}» if none is coming"))

        # the applied column against the status
        if cfg.has("applied") and status:
            flag = rec.get("applied")
            says_yes = flag is True or str(flag or "").strip().casefold() == cfg.values.yes.casefold()
            says_no = flag is False or str(flag or "").strip().casefold() == cfg.values.no.casefold()
            if says_no and (status in t.waiting or t._stage(status) is not None):
                out.append(Issue("applied_mismatch", r, "applied", flag,
                                 f"says not applied, but the status is «{status}»"))
            elif says_yes and status == st.not_applied:
                out.append(Issue("applied_mismatch", r, "applied", flag,
                                 f"says applied, but the status is «{status}»"))

        # the furthest stage
        if cfg.has("reached") and status:
            cur = t._stage(status)
            old_raw = ws.cell(r, lay.cols["reached"]).value
            old = t._stage(old_raw)
            if cur is not None and (old is not None or old_raw in (None, "")) and (old is None or cur > old):
                new = t._stage_word(ws, lay, cur)
                out.append(Issue("reached_behind", r, "reached", old_raw,
                                 f"the status «{status}» is further than «{old_raw or '∅'}»", fix=new))

        # the reply date
        if cfg.has("replied_on") and t._furthest({**rec, "status": status}) >= 0 \
                and ws.cell(r, lay.cols["replied_on"]).value in (None, ""):
            d, why = reply_date_hint(t, rec, audit_days.get(r))
            out.append(Issue("no_reply_date", r, "replied_on", None,
                             f"replied ({status}); {why}" if d else
                             f"replied ({status}) but no date found; set it with set_status(on=...) or by hand",
                             fix=d if d else NOFIX))

        # duplicates
        key = (company_key(rec.get("company")), str(rec.get("role") or "").strip().casefold())
        if key[0] and key in seen:
            out.append(Issue("duplicate", r, None, None, f"the same company and role as row {seen[key]}"))
        else:
            seen.setdefault(key, r)
    return out


def check(t: Tracker, limit: int = 200) -> dict[str, Any]:
    _, ws, lay, version = t._load()
    issues = find_issues(t, ws, lay)
    by_kind: dict[str, dict[str, Any]] = {}
    for i in issues:
        k = by_kind.setdefault(i.kind, {"count": 0, "fixable": 0, "what": KINDS[i.kind]})
        k["count"] += 1
        k["fixable"] += i.fixable
    shown = [{"kind": i.kind, "row": i.row, "label": t._label(ws, lay, i.row),
              **({"field": i.field} if i.field else {}),
              **({"value": t._show(i.value, i.field or "")} if i.value is not None else {}),
              "note": i.note,
              **({"fix": t._show(i.fix, i.field or "")} if i.fixable else {})}
             for i in issues[:limit]]
    fixable = sum(i.fixable for i in issues)
    return {"count": len(issues), "fixable": fixable, "by_kind": by_kind, "issues": shown,
            "truncated": len(issues) > limit,
            "next": ("fix_tracker writes the fixable ones in one change, with one backup and one undo "
                     "(kinds=[...] or rows=[...] to choose); the rest need a decision.") if fixable else
                    "Nothing the server can fix by itself.",
            "version": version}


def fix(t: Tracker, kinds: list[str] | None = None, rows: list[int] | None = None, *,
        expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
    unknown = sorted(set(kinds or []) - set(KINDS))
    if unknown:
        raise TrackerError(f"Unknown kind(s) {', '.join(unknown)}. Kinds: {', '.join(KINDS)}")

    def mutate(wb, ws, lay):
        picked = [i for i in find_issues(t, ws, lay) if i.fixable
                  and (kinds is None or i.kind in kinds) and (rows is None or i.row in rows)]
        # one value per cell: a later check wins (a text date on a Sunday becomes Monday's date)
        values = {(i.row, i.field): i.fix for i in picked}
        changes = []
        for (row, field), value in values.items():
            t._set(ws, lay, row, field, value, changes)
        return sorted({row for row, _ in values}), changes, []

    return t._write("fix_tracker", mutate, expected_version, dry_run)

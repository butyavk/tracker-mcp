"""A tracker kept in Google Sheets.

The target is the sheet's URL (`https://docs.google.com/spreadsheets/d/<id>/...`)
or `gsheet:<id>`. The store reads every sheet through the Sheets API into an
in-memory workbook, so the tracker's logic is the same as for an .xlsx file, and
writes back only the cells that changed, in one atomic `batchUpdate`.

Safety, compared with a file:
- version = hash of all cell values; a write is refused if it moved since the read
  (the API has no If-Match for values, so a short window remains, as with files);
- strings are written as strings (`stringValue`), never parsed: a note starting
  with "=" stays text and cannot become a formula;
- each write keeps a local JSON backup of the old and new value of every cell it
  touched; undo writes the old values back;
- after writing, the cells are read back and checked; on a mismatch the old values
  are restored and the write is reported as failed.

Authorization, either of:
- your own Google account: `tracker-mcp auth google --client-secret client_secret.json`
  once (an OAuth "Desktop app" client from Google Cloud); the refresh token is kept
  in ~/.tracker-mcp/google-token.json;
- a service account: GOOGLE_APPLICATION_CREDENTIALS=key.json, and share the sheet
  with the service account's e-mail.
Scope: https://www.googleapis.com/auth/spreadsheets only. No Google libraries are
needed: plain HTTPS (urllib) plus PyJWT for the service-account assertion.
"""

from __future__ import annotations

import base64
import hashlib
import http.server
import json
import os
import re
import secrets
import time
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from collections.abc import Callable
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .stores import ConflictError, TrackerError, Verify, digest

SCOPE = "https://www.googleapis.com/auth/spreadsheets"
API = "https://sheets.googleapis.com/v4/spreadsheets"
TOKEN_URI = "https://oauth2.googleapis.com/token"
EPOCH = datetime(1899, 12, 30)
FIELDS = ("properties.title,sheets(properties(sheetId,title,gridProperties(rowCount,columnCount)),"
          "data(rowData(values(userEnteredValue,effectiveFormat.numberFormat,dataValidation))))")

# A transport takes (method, url, json body or None) and returns the parsed JSON.
Transport = Callable[[str, str, dict | None], dict]


def home_dir() -> Path:
    return Path(os.environ.get("TRACKER_MCP_HOME") or Path.home() / ".tracker-mcp").expanduser()


def sheet_id(target: str) -> str:
    if target.startswith("gsheet:"):
        return target.split(":", 1)[1].strip("/")
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]+)", target)
    if not m:
        raise TrackerError(f"Not a Google Sheets URL: {target}")
    return m.group(1)


def default_config_path(target: str) -> Path:
    return home_dir() / f"{sheet_id(target)}.tracker-mcp.toml"


# ---------------------------------------------------------------- auth

class GoogleAuth:
    """Access tokens from a service-account key or from the stored user token."""

    def __init__(self, token_file: Path | None = None):
        self.token_file = token_file or Path(os.environ.get("TRACKER_GOOGLE_TOKEN") or home_dir() / "google-token.json")
        self._token: tuple[str, float] | None = None

    def header(self) -> str:
        if not self._token or self._token[1] < time.time() + 60:
            self._token = self._fetch()
        return f"Bearer {self._token[0]}"

    def reset(self) -> None:
        self._token = None

    def _fetch(self) -> tuple[str, float]:
        key_file = os.environ.get("GOOGLE_APPLICATION_CREDENTIALS")
        if key_file:
            info = json.loads(Path(key_file).expanduser().read_text(encoding="utf-8"))
            if info.get("type") == "service_account":
                return self._service_account(info)
        if self.token_file.exists():
            info = json.loads(self.token_file.read_text(encoding="utf-8"))
            return self._refresh(info)
        raise TrackerError("Not signed in to Google. Run `tracker-mcp auth google --client-secret "
                           "<client_secret.json>` once, or set GOOGLE_APPLICATION_CREDENTIALS to a "
                           "service-account key and share the sheet with its e-mail.")

    def _service_account(self, info: dict) -> tuple[str, float]:
        import jwt  # PyJWT, already a dependency of the MCP SDK

        now = int(time.time())
        claims = {"iss": info["client_email"], "scope": SCOPE, "aud": info.get("token_uri", TOKEN_URI),
                  "iat": now, "exp": now + 3600}
        assertion = jwt.encode(claims, info["private_key"], algorithm="RS256")
        data = _post_form(info.get("token_uri", TOKEN_URI),
                          {"grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer", "assertion": assertion})
        return data["access_token"], time.time() + int(data.get("expires_in", 3600))

    @staticmethod
    def _refresh(info: dict) -> tuple[str, float]:
        data = _post_form(info.get("token_uri", TOKEN_URI), {
            "grant_type": "refresh_token", "refresh_token": info["refresh_token"],
            "client_id": info["client_id"], "client_secret": info.get("client_secret", "")})
        return data["access_token"], time.time() + int(data.get("expires_in", 3600))


def _post_form(url: str, fields: dict[str, str]) -> dict:
    req = urllib.request.Request(url, data=urllib.parse.urlencode(fields).encode(), method="POST",
                                 headers={"Content-Type": "application/x-www-form-urlencoded"})
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            return json.loads(resp.read())
    except urllib.error.HTTPError as e:
        detail = e.read().decode(errors="replace")[:300]
        raise TrackerError(f"Google sign-in failed ({e.code}): {detail}. Run `tracker-mcp auth google` "
                           "again if the token was revoked.") from e


def login(client_secret: str | Path, token_file: Path | None = None, open_browser: bool = True) -> Path:
    """One-time OAuth sign-in for a Desktop-app client (loopback redirect + PKCE)."""
    raw = json.loads(Path(client_secret).expanduser().read_text(encoding="utf-8"))
    client = raw.get("installed") or raw.get("web") or raw
    verifier = secrets.token_urlsafe(64)
    challenge = base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b"=").decode()
    state = secrets.token_urlsafe(16)
    result: dict[str, str] = {}

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            q = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(self.path).query))
            result.update(q)
            self.send_response(200)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.end_headers()
            self.wfile.write(b"tracker-mcp: signed in. You can close this tab.")

        def log_message(self, *args):
            pass

    server = http.server.HTTPServer(("127.0.0.1", 0), Handler)
    redirect = f"http://127.0.0.1:{server.server_port}/"
    url = client.get("auth_uri", "https://accounts.google.com/o/oauth2/auth") + "?" + urllib.parse.urlencode({
        "client_id": client["client_id"], "redirect_uri": redirect, "response_type": "code", "scope": SCOPE,
        "code_challenge": challenge, "code_challenge_method": "S256", "access_type": "offline",
        "prompt": "consent", "state": state})
    print(f"Open this link to sign in to Google:\n{url}\n", flush=True)
    if open_browser:
        webbrowser.open(url)
    while "code" not in result and "error" not in result:
        server.handle_request()
    server.server_close()
    if result.get("state") != state or "code" not in result:
        raise TrackerError(f"Google sign-in did not complete: {result.get('error', 'state mismatch')}")
    token_uri = client.get("token_uri", TOKEN_URI)
    data = _post_form(token_uri, {"grant_type": "authorization_code", "code": result["code"],
                                  "redirect_uri": redirect, "client_id": client["client_id"],
                                  "client_secret": client.get("client_secret", ""), "code_verifier": verifier})
    if "refresh_token" not in data:
        raise TrackerError("Google returned no refresh token. Remove tracker-mcp from your Google account's "
                           "third-party access and sign in again.")
    target = token_file or GoogleAuth().token_file
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps({"client_id": client["client_id"], "client_secret": client.get("client_secret", ""),
                                  "refresh_token": data["refresh_token"], "token_uri": token_uri,
                                  "scope": SCOPE}, indent=1), encoding="utf-8")
    try:
        target.chmod(0o600)
    except OSError:
        pass
    return target


def http_transport(auth: GoogleAuth) -> Transport:
    def call(method: str, url: str, body: dict | None) -> dict:
        for attempt in (1, 2):
            req = urllib.request.Request(url, method=method, data=json.dumps(body).encode() if body else None,
                                         headers={"Authorization": auth.header(),
                                                  "Content-Type": "application/json"})
            try:
                with urllib.request.urlopen(req, timeout=60) as resp:
                    return json.loads(resp.read() or b"{}")
            except urllib.error.HTTPError as e:
                if e.code == 401 and attempt == 1:
                    auth.reset()
                    continue
                raise TrackerError(_api_error(e)) from e
        raise TrackerError("Google Sheets: authorization keeps failing")
    return call


def _api_error(e: urllib.error.HTTPError) -> str:
    try:
        msg = json.loads(e.read()).get("error", {}).get("message", "")
    except (ValueError, AttributeError):
        msg = ""
    hint = {403: " Share the sheet with the account you signed in with (or the service account's e-mail).",
            404: " Check the sheet URL.",
            429: " Too many requests; retry in a minute."}.get(e.code, "")
    return f"Google Sheets error {e.code}: {msg}.{hint}"


# ---------------------------------------------------------------- store

def _from_api(cell: dict) -> Any:
    v = cell.get("userEnteredValue")
    if not v:
        return None
    if "formulaValue" in v:
        return v["formulaValue"]
    if "stringValue" in v:
        return v["stringValue"]
    if "boolValue" in v:
        return v["boolValue"]
    if "numberValue" in v:
        n = v["numberValue"]
        kind = cell.get("effectiveFormat", {}).get("numberFormat", {}).get("type")
        if kind in ("DATE", "DATE_TIME"):
            return EPOCH + timedelta(days=n)
        return int(n) if isinstance(n, float) and n.is_integer() else n
    return None


def _to_api(value: Any, was_formula: bool) -> tuple[dict, str]:
    """CellData and its field mask. Text never becomes a formula unless the cell held one."""
    if value is None or value == "":
        return {}, "userEnteredValue"
    if isinstance(value, bool):
        return {"userEnteredValue": {"boolValue": value}}, "userEnteredValue"
    if isinstance(value, (datetime, date)):
        d = value if isinstance(value, datetime) else datetime(value.year, value.month, value.day)
        serial = (d - EPOCH).total_seconds() / 86400
        return ({"userEnteredValue": {"numberValue": serial},
                 "userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "yyyy-mm-dd"}}},
                "userEnteredValue,userEnteredFormat.numberFormat")
    if isinstance(value, (int, float)):
        return {"userEnteredValue": {"numberValue": value}}, "userEnteredValue"
    text = str(value)
    if was_formula and text.startswith("="):
        return {"userEnteredValue": {"formulaValue": text}}, "userEnteredValue"
    return {"userEnteredValue": {"stringValue": text}}, "userEnteredValue"


def _plain(value: Any) -> Any:
    """A JSON-friendly, comparable form of a cell value."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


class SheetsStore:
    kind = "gsheet"
    single_sheet = False

    def __init__(self, target: str, backup_dir: str | Path | None = None, transport: Transport | None = None):
        self.id = sheet_id(target)
        self.path = None
        self.label = f"Google Sheet {self.id}"
        self.backup_dir = Path(backup_dir).expanduser() if backup_dir else home_dir() / "backups" / self.id
        self._call = transport or http_transport(GoogleAuth())
        self._loaded: dict[str, dict] = {}   # version -> {"cells", "sheets"} of that read

    # ---- reading

    def _fetch(self) -> dict:
        return self._call("GET", f"{API}/{self.id}?includeGridData=true&fields={urllib.parse.quote(FIELDS)}", None)

    def _read(self) -> tuple[Workbook, str, dict]:
        doc = self._fetch()
        wb = Workbook()
        wb.remove(wb.active)
        cells: dict[tuple[str, int, int], Any] = {}
        sheets: dict[str, dict] = {}
        for sheet in doc.get("sheets", []):
            props = sheet["properties"]
            title = props["title"]
            ws = wb.create_sheet(title)
            sheets[title] = {"id": props["sheetId"], "rows": props.get("gridProperties", {}).get("rowCount", 1000),
                             "cols": props.get("gridProperties", {}).get("columnCount", 26)}
            lists: dict[int, list[str]] = {}
            ranges: dict[int, str] = {}
            for block in sheet.get("data", []):
                for r, row in enumerate(block.get("rowData", []), start=1):
                    for c, cell in enumerate(row.get("values", []), start=1):
                        value = _from_api(cell)
                        if value is not None:
                            ws.cell(r, c, value)
                            cells[(title, r, c)] = value
                        cond = cell.get("dataValidation", {}).get("condition", {})
                        if cond.get("type") == "ONE_OF_LIST" and c not in lists:
                            lists[c] = [v.get("userEnteredValue", "") for v in cond.get("values", [])]
                        elif cond.get("type") == "ONE_OF_RANGE" and c not in ranges:
                            ranges[c] = cond.get("values", [{}])[0].get("userEnteredValue", "").lstrip("=")
            ws.tracker_dropdowns = lists            # read by tracker.dropdown_values
            last = max(ws.max_row, 2)
            for c, ref in ranges.items():
                dv = DataValidation(type="list", formula1=ref)
                dv.add(f"{get_column_letter(c)}1:{get_column_letter(c)}{last}")
                ws.add_data_validation(dv)
        snapshot = {"|".join(map(str, k)): _plain(v) for k, v in sorted(cells.items())}
        version = digest(json.dumps(snapshot, sort_keys=True, ensure_ascii=False, default=str).encode())
        self._loaded = {version: {"cells": cells, "sheets": sheets}}
        return wb, version, {"cells": cells, "sheets": sheets}

    def load(self) -> tuple[Workbook, str, list[str]]:
        wb, version, _ = self._read()
        return wb, version, []

    def version(self) -> str:
        return self._read()[1]

    def check_writable(self) -> None:
        return None   # the API has no lock; conflicts are caught by the version check

    # ---- writing

    def _diff(self, wb: Workbook, base: dict) -> list[dict]:
        out = []
        for title in base["sheets"]:
            if title not in wb.sheetnames:
                continue
            ws = wb[title]
            seen = set()
            for row in ws.iter_rows():
                for cell in row:
                    key = (title, cell.row, cell.column)
                    seen.add(key)
                    old = base["cells"].get(key)
                    if _plain(cell.value) != _plain(old):
                        out.append({"sheet": title, "row": cell.row, "col": cell.column,
                                    "old": old, "new": cell.value})
            for key, old in base["cells"].items():
                if key[0] == title and key not in seen:
                    out.append({"sheet": title, "row": key[1], "col": key[2], "old": old, "new": None})
        return out

    def _requests(self, changes: list[dict], sheets: dict, value_key: str) -> list[dict]:
        reqs: list[dict] = []
        for title, meta in sheets.items():
            need = max((ch["row"] for ch in changes if ch["sheet"] == title), default=0)
            if need > meta["rows"]:
                reqs.append({"appendDimension": {"sheetId": meta["id"], "dimension": "ROWS",
                                                 "length": need - meta["rows"] + 50}})
        for ch in changes:
            old = ch["old"] if value_key == "new" else ch["new"]
            data, fields = _to_api(ch[value_key], isinstance(old, str) and old.startswith("="))
            reqs.append({"updateCells": {
                "range": {"sheetId": sheets[ch["sheet"]]["id"], "startRowIndex": ch["row"] - 1,
                          "endRowIndex": ch["row"], "startColumnIndex": ch["col"] - 1, "endColumnIndex": ch["col"]},
                "rows": [{"values": [data]}], "fields": fields}})
        return reqs

    def _backup(self, changes: list[dict], before: str) -> Path:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        path = self.backup_dir / f"{self.id}.{stamp}.json"
        path.write_text(json.dumps({"spreadsheet": self.id, "version_before": before,
                                    "cells": [{**c, "old": _plain(c["old"]), "new": _plain(c["new"])}
                                              for c in changes]},
                                   ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        return path

    def _apply(self, changes: list[dict], sheets: dict, value_key: str) -> None:
        if changes:
            self._call("POST", f"{API}/{self.id}:batchUpdate",
                       {"requests": self._requests(changes, sheets, value_key)})

    def save(self, wb: Workbook, before: str, verify: Verify) -> tuple[str, str]:
        base = self._loaded.get(before)
        current_wb, now, current = self._read()
        if now != before or base is None:
            raise ConflictError("The sheet changed since it was read; nothing was saved. Re-read and retry.")
        changes = self._diff(wb, current)
        backup = self._backup(changes, before)
        self._apply(changes, current["sheets"], "new")
        written, after, _ = self._read()
        try:
            verify(written)
        except TrackerError:
            self._apply(changes, current["sheets"], "old")   # put the old values back
            raise
        return backup.name, after

    def restore(self, backup: str) -> tuple[str, str]:
        path = self.backup_dir / backup
        if not path.exists():
            raise TrackerError(f"Backup {backup} is gone; nothing to restore.")
        saved = json.loads(path.read_text(encoding="utf-8"))
        _, now, current = self._read()
        changes = [{**c, "old": c["new"], "new": c["old"]} for c in saved["cells"]]
        for c in changes:   # dates come back from JSON as text; write them as dates again
            if isinstance(c["new"], str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", c["new"]):
                c["new"] = datetime.fromisoformat(c["new"])
        safety = self._backup(changes, now)
        self._apply(changes, current["sheets"], "new")
        return safety.name, self._read()[1]

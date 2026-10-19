"""Where a tracker lives: an .xlsx file, a .csv file or a Google Sheet.

A store hands the tracker an in-memory workbook (openpyxl) together with a
version, and writes a changed workbook back. The tracker's logic only ever sees
the workbook, so every store gets the same tools and the same safety rules:

- `load()` returns the workbook, its version and any warnings about parts of
  the source that a save would lose;
- `save()` refuses if the version moved since `load()`, makes a backup, writes,
  and lets the tracker verify what was written before the write counts;
- `restore()` puts a backup back (undo).

The Google Sheets store lives in gsheets.py; `open_store` picks a store by target.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import warnings
import zipfile
from collections.abc import Callable
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Protocol

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

KEEP_BACKUPS = 200


class TrackerError(Exception):
    """A request the tracker cannot apply; the message says why."""


class ConflictError(TrackerError):
    """The source changed (or is open elsewhere) under us; re-read and retry."""


class UnverifiedWrite(TrackerError):
    """The write was sent but could not be read back to check it."""

    def __init__(self, backup: str, message: str):
        super().__init__(message)
        self.backup = backup


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


Verify = Callable[[Workbook], None]


class Store(Protocol):
    kind: str                  # "xlsx", "csv" or "gsheet"
    label: str                 # how messages name the tracker
    backup_dir: Path           # backups and audit.jsonl
    single_sheet: bool         # True when the source has exactly one sheet (CSV)

    def load(self) -> tuple[Workbook, str, list[str]]: ...
    def version(self) -> str: ...
    def check_writable(self) -> None: ...
    def save(self, wb: Workbook, before: str, verify: Verify) -> tuple[str, str]: ...
    def restore(self, backup: str) -> tuple[str, str]: ...


def is_google_sheet(target: str) -> bool:
    return target.startswith(("https://docs.google.com/spreadsheets/", "gsheet:"))


def open_store(target: str | Path, backup_dir: str | Path | None = None) -> Store:
    t = str(target)
    if is_google_sheet(t):
        from .gsheets import SheetsStore
        return SheetsStore(t, backup_dir)
    suffix = Path(t).suffix.lower()
    if suffix == ".csv":
        return CsvStore(t, backup_dir)
    if suffix == ".xlsx":
        return XlsxStore(t, backup_dir)
    # .xlsm would lose its macros and its file type on save; .xls and .ods are other formats.
    raise TrackerError(f"tracker-mcp works with .xlsx and .csv files and Google Sheets, not «{suffix or t}». "
                       "Save a copy of the tracker as .xlsx (Excel or Numbers: File → Save As / Export) "
                       "and point the server to it.")


# ---------------------------------------------------------------- files

class _FileStore:
    """What .xlsx and .csv share: a local file, lock files, backups next to it,
    and the write path temp file -> verify -> version check -> backup -> rename."""

    kind = "file"
    single_sheet = False

    def __init__(self, path: str | Path, backup_dir: str | Path | None = None):
        self.path = Path(path).expanduser()
        self.label = str(self.path)
        self.backup_dir = (Path(backup_dir).expanduser() if backup_dir
                           else self.path.parent / "backups-tracker-mcp")

    def _bytes(self) -> bytes:
        if not self.path.exists():
            raise TrackerError(f"Tracker file not found: {self.path}")
        return self.path.read_bytes()

    def version(self) -> str:
        return digest(self._bytes())

    def check_writable(self) -> None:
        for lock in (f"~${self.path.name}", f".~lock.{self.path.name}#"):
            if (self.path.parent / lock).exists():
                raise ConflictError(f"The tracker is open in Excel/LibreOffice (lock file {lock}). "
                                    "Close it and retry; if it is closed, the lock file is stale.")

    def _render(self, wb: Workbook) -> bytes:
        raise NotImplementedError

    def _parse(self, data: bytes) -> Workbook:
        raise NotImplementedError

    def _tmp(self) -> Path:
        return self.path.with_name(f".{self.path.stem}.tracker-mcp-tmp{self.path.suffix}")

    def _replace(self, tmp: Path) -> None:
        try:
            os.replace(tmp, self.path)
        except OSError as e:  # Windows: OneDrive, an antivirus or Excel holding the file
            raise ConflictError(f"Could not replace {self.path.name}: {e.strerror or e}. Another program "
                                "(Excel, OneDrive, an antivirus) may be holding it. Nothing was changed; "
                                "retry in a moment.") from e

    def save(self, wb: Workbook, before: str, verify: Verify) -> tuple[str, str]:
        tmp = self._tmp()
        try:
            tmp.write_bytes(self._render(wb))
            verify(self._parse(tmp.read_bytes()))
            if self.version() != before:  # someone saved while we worked
                raise ConflictError("The tracker changed while writing; nothing was saved. Retry.")
            backup = self._backup()
            shutil.copymode(self.path, tmp)
            self._replace(tmp)
        finally:
            tmp.unlink(missing_ok=True)
        return backup.name, self.version()

    def restore(self, backup: str) -> tuple[str, str]:
        self.check_writable()
        source = self.backup_dir / backup
        if not source.exists():
            raise TrackerError(f"Backup {backup} is gone; nothing to restore.")
        safety = self._backup()
        tmp = self._tmp()
        try:
            shutil.copy2(source, tmp)
            self._replace(tmp)
        finally:
            tmp.unlink(missing_ok=True)
        return safety.name, self.version()

    def _backup(self) -> Path:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        dst = self.backup_dir / f"{self.path.stem}.{stamp}{self.path.suffix}"
        shutil.copy2(self.path, dst)
        old = sorted(self.backup_dir.glob(f"{self.path.stem}.*{self.path.suffix}"))
        for p in old[:-KEEP_BACKUPS]:
            p.unlink()
        return dst


class XlsxStore(_FileStore):
    kind = "xlsx"

    def _parse(self, data: bytes) -> Workbook:
        try:
            return openpyxl.load_workbook(BytesIO(data))
        except (zipfile.BadZipFile, InvalidFileException, KeyError) as e:
            raise TrackerError(f"{self.path.name} is not a readable .xlsx file ({e}). If it was saved "
                               "in another format, save it again as an Excel Workbook (.xlsx).") from e

    def load(self) -> tuple[Workbook, str, list[str]]:
        data = self._bytes()
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            wb = self._parse(data)
        # Parts of the file openpyxl cannot keep (e.g. dropdowns that newer Excel stores
        # in an extension): a save would silently drop them, so writes are refused.
        lossy = sorted({str(w.message) for w in caught if "will be removed" in str(w.message)})
        return wb, digest(data), lossy

    def _render(self, wb: Workbook) -> bytes:
        wb.calculation.fullCalcOnLoad = True  # Excel recomputes formulas on open
        out = BytesIO()
        wb.save(out)
        return out.getvalue()


class CsvStore(_FileStore):
    """A CSV export (Notion, Huntr, Teal, Google Sheets...). One sheet, no dropdowns:
    statuses come from the config. Dates are written as yyyy-mm-dd text."""

    kind = "csv"
    single_sheet = True

    def __init__(self, path: str | Path, backup_dir: str | Path | None = None):
        super().__init__(path, backup_dir)
        self.delimiter, self.bom, self.encoding, self.newline = ",", False, "utf-8", "\r\n"
        self.sheet = re.sub(r"[\[\]:*?/\\]", "_", self.path.stem)[:31] or "Sheet1"

    def _parse(self, data: bytes) -> Workbook:
        self.bom = data.startswith(b"\xef\xbb\xbf")
        try:
            text, self.encoding = data.decode("utf-8-sig"), "utf-8"
        except UnicodeDecodeError:   # Excel's plain "CSV" on Windows
            text, self.encoding = data.decode("cp1252", errors="replace"), "cp1252"
        self.newline = "\r\n" if "\r\n" in text[:65536] or "\n" not in text else "\n"
        try:
            self.delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            self.delimiter = ","
        wb = Workbook()
        ws = wb.active
        ws.title = self.sheet
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=self.delimiter):
            ws.append([ILLEGAL_CHARACTERS_RE.sub(" ", v) if v != "" else None for v in row])
        return wb

    def load(self) -> tuple[Workbook, str, list[str]]:
        data = self._bytes()
        return self._parse(data), digest(data), []

    def _render(self, wb: Workbook) -> bytes:
        out = io.StringIO(newline="")
        writer = csv.writer(out, delimiter=self.delimiter, lineterminator=self.newline)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            writer.writerow([_csv_text(v) for v in row])
        try:
            body = out.getvalue().encode(self.encoding)
        except UnicodeEncodeError as e:
            raise TrackerError(f"«{e.object[e.start:e.end]}» cannot be written to this CSV file: it is saved "
                               f"as {self.encoding}. Save it as «CSV UTF-8» in Excel and retry.") from e
        return (b"\xef\xbb\xbf" if self.bom else b"") + body


def _csv_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)

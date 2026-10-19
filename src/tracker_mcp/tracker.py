"""Read and safely write a job-application tracker: .xlsx, .csv or Google Sheets.

The tracker stays the source of truth: people open it in Excel or in the browser,
and other scripts edit it too. So the server never caches it. Every call re-reads
it through its store (stores.py, gsheets.py), and every write goes through
`Tracker._write`, which:

1. refuses while the file is open in Excel or LibreOffice (lock file present);
2. refuses if the tracker changed since the caller last read it: `expected_version`
   works like an HTTP If-Match, the version is a hash of the content;
3. writes, re-reads and checks every intended cell (a file: in a temp copy, before
   it replaces the original);
4. keeps a backup, so `undo_last_write` can put the old state back;
5. appends an audit record (field, old -> new, backup name).

Which sheet, which columns, which status words and which follow-up rule: all of
that comes from the config (config.py). Columns are found by their header text,
not by letter, so inserting a column in Excel does not break the server.
"""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from collections.abc import Callable
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.worksheet import Worksheet

from .config import Config, load_config_file, parse_config, with_column
from .stores import ConflictError, Store, TrackerError, UnverifiedWrite, backup_path, open_store

__all__ = ["ConflictError", "Tracker", "TrackerError", "dropdown_values", "parse_date"]

EXCEL_EPOCH = date(1899, 12, 30)


# ---------------------------------------------------------------- values

def parse_date(value: Any, formats: tuple[str, ...] | list[str] = ()) -> date | None:
    """Dates in a hand-kept sheet are mixed: datetimes, ISO strings, Excel serials,
    dd.mm.yyyy, and whatever `formats` (strptime patterns from the config) add."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 30000 < value < 80000:
        return EXCEL_EPOCH + timedelta(days=int(value))
    if isinstance(value, str):
        s = value.strip()
        if s.isdigit() and 30000 < int(s) < 80000:
            return EXCEL_EPOCH + timedelta(days=int(s))
        try:
            if m := re.match(r"(\d{4})-(\d{2})-(\d{2})", s):
                return date(*map(int, m.groups()))
            if m := re.match(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", s):
                d, mo, y = map(int, m.groups())
                return date(y, mo, d)
        except ValueError:
            return None
        for fmt in formats:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def sheet_in(wb: Workbook, title: str) -> str:
    """The workbook's name for a sheet title. A Google tab title openpyxl refuses
    ("Q1/Q2") gets another name in the workbook; configs keep the real title."""
    return getattr(wb, "tracker_titles", {}).get(title, title)


def real_title(ws: Worksheet) -> str:
    """The sheet's title as the user sees it (the Google tab title)."""
    back = {safe: real for real, safe in getattr(ws.parent, "tracker_titles", {}).items()}
    return back.get(ws.title, ws.title)


def dropdown_values(ws: Worksheet, col: int) -> list[str] | None:
    """The list a column's dropdown offers: inline ("A,B,C"), a cell range
    (Lists!$A$1:$A$9) or a defined name pointing to one. None if there is none."""
    lists = getattr(ws, "tracker_dropdowns", None)   # set by the Google Sheets store
    if lists and lists.get(col):
        return [s.strip() for s in lists[col] if str(s).strip()]
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1 or not any(
                rng.min_col <= col <= rng.max_col for rng in dv.sqref.ranges):
            continue
        f = dv.formula1.strip().lstrip("=")
        if f.startswith('"'):
            return [s.strip() for s in f.strip('"').split(",") if s.strip()]
        wb = ws.parent
        if f in wb.defined_names:
            f = wb.defined_names[f].attr_text
        sheet_name, _, ref = f.rpartition("!")
        sheet_name = sheet_in(wb, sheet_name.strip("'"))
        try:
            sheet = wb[sheet_name] if sheet_name else ws
            cells = sheet[ref.replace("$", "")]
        except (KeyError, ValueError):
            return None
        flat = [c for row in (cells if isinstance(cells, tuple) else ((cells,),)) for c in
                (row if isinstance(row, tuple) else (row,))]
        return [str(c.value).strip() for c in flat if c.value not in (None, "")] or None
    return None


def weekday_or_monday(d: date) -> date:
    return d + timedelta(days=(7 - d.weekday()) % 7) if d.weekday() >= 5 else d


def company_key(name: Any) -> str:
    """A company name for comparing: "Swap (Swap Commerce)", "SWAP Inc." -> "swap"."""
    s = re.sub(r"\(.*?\)|\[.*?\]", " ", str(name or "")).casefold()
    s = re.sub(r"\b(inc|llc|ltd|limited|gmbh|corp|corporation|plc|ag|bv|oy|ab|group|holdings?)\b\.?", " ", s)
    return re.sub(r"[\W_]+", "", s)


_LEAD_DATE = re.compile(r"^\W{0,3}(\d{4}-\d{2}-\d{2}|\d{1,2}\.\d{1,2}(?:\.\d{2,4})?)(?!\d)")
_ANY_DATE = re.compile(r"(?<![\w.])(\d{4}-\d{2}-\d{2}|\d{1,2}\.\d{1,2}(?:\.\d{4})?)(?![\d])")


def date_in_text(token: str, not_after: date, not_before: date | None = None) -> date | None:
    """A date written in a note: 2026-10-04, 04.10.2026, or 04.10 without a year (the
    latest such date that is not after `not_after`, and not before `not_before`)."""
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", token):
            d = date.fromisoformat(token)
            return d if d <= not_after and (not_before is None or d >= not_before) else None
        parts = [int(p) for p in token.split(".")]
        if len(parts) == 3:
            y = parts[2] + 2000 if parts[2] < 100 else parts[2]
            d = date(y, parts[1], parts[0])
            return d if d <= not_after and (not_before is None or d >= not_before) else None
        for year in (not_after.year, not_after.year - 1):
            d = date(year, parts[1], parts[0])
            if d <= not_after and (not_before is None or d >= not_before):
                return d
    except ValueError:
        return None
    return None


@dataclass
class Change:
    row: int
    field: str
    old: Any
    new: Any  # the raw value written: str, int or date

    def as_dict(self, show: Callable[[Any, str], Any]) -> dict[str, Any]:
        return {"row": self.row, "field": self.field, "old": self.old, "new": show(self.new, self.field)}


@dataclass
class Layout:
    header_row: int
    cols: dict[str, int]       # field -> column number
    headers: dict[int, str]    # column number -> header text
    last_row: int              # last row with a company name

    def rows(self) -> range:
        return range(self.header_row + 1, self.last_row + 1)


# ---------------------------------------------------------------- tracker


def undid(undo: dict[str, Any], entry: dict[str, Any]) -> bool:
    """Whether an undo entry of the audit log undid this write. The backup name is unique
    per write; the time (to the second) is what older logs have."""
    if undo.get("undid_backup"):
        return undo["undid_backup"] == entry.get("backup")
    return undo.get("undid") == entry.get("at")


class Tracker:
    def __init__(self, target: str | Path, *, config: Config, backup_dir: str | Path | None = None,
                 today: Callable[[], date] = date.today, store: Store | None = None,
                 config_file: str | Path | None = None):
        self.store = store or open_store(target, backup_dir)
        self.path: Path | None = getattr(self.store, "path", None)
        self.backup_dir = self.store.backup_dir
        # where the config lives, for add_column(field=...); None: the server does not write it
        self.config_file = Path(config_file).expanduser() if config_file else None
        self.today = today
        self.lossy: list[str] = []   # set by each load, see _load
        self._use(config)

    def _use(self, config: Config) -> None:
        self.cfg = config
        st = config.statuses
        self.closed = set(st.closed)
        self.waiting = set(st.waiting or [st.applied])
        self.rejected = set(st.rejected)
        self.aliases = {k.casefold(): v for k, v in st.aliases.items()}
        # funnel stage of a status (or of a stage name, for the `reached` column)
        self.stage_names = list(st.stages)
        self.stage_of: dict[str, int] = {}
        for i, (name, members) in enumerate(st.stages.items()):
            for s in (name, *members):
                self.stage_of[s.casefold()] = i

    # ---- values

    def _date(self, value: Any) -> date | None:
        return parse_date(value, self.cfg.values.date_formats)

    def _show(self, value: Any, key: str = "") -> Any:
        """Cell value -> JSON-friendly value."""
        if self.cfg.is_date(key):
            d = self._date(value)
            if d:
                return d.isoformat()
        if isinstance(value, (datetime, date)):
            return (value.date() if isinstance(value, datetime) else value).isoformat()
        if isinstance(value, str):
            return value.strip() or None
        return value

    def _append(self, old: Any, text: str) -> str:
        old = str(old).strip() if old is not None else ""
        return f"{old}{self.cfg.values.separator}{text}" if old else text

    def _stamp(self, text: str, on: date | None = None) -> str:
        return f"{(on or self.today()).strftime(self.cfg.values.stamp_format)}: {text.strip()}"

    def _fmt(self, d: date) -> str:
        return d.strftime(self.cfg.values.stamp_format)

    def _next_date(self, start: date, days: int) -> date:
        d = start + timedelta(days=days)
        return weekday_or_monday(d) if self.cfg.followups.weekend_to_monday else d

    def _is_applied(self, rec: dict[str, Any]) -> bool:
        if self.cfg.has("applied"):
            v = rec.get("applied")
            if isinstance(v, bool):  # a checkbox column
                return v
            return v is not None and str(v).strip().casefold() == self.cfg.values.yes.casefold()
        return rec.get("status") not in (None, self.cfg.statuses.not_applied)

    def _flag(self, ws: Worksheet, lay: Layout, yes: bool) -> bool | str:
        """The value for the applied column: a bool if the column holds checkboxes."""
        col = lay.cols["applied"]
        if any(isinstance(ws.cell(r, col).value, bool) for r in lay.rows()):
            return yes
        return self.cfg.values.yes if yes else self.cfg.values.no

    def _is_reply(self, status: Any) -> bool:
        """A status that means the company answered: a rejection or any funnel stage."""
        return status in self.rejected or self._stage(status) is not None

    def _stage(self, value: Any) -> int | None:
        return self.stage_of.get(str(value).strip().casefold()) if value not in (None, "") else None

    def _stage_word(self, ws: Worksheet, lay: Layout, stage: int) -> str:
        """What to write into the reached column: its own dropdown word for the stage
        if it has one, else the stage's first status."""
        for option in dropdown_values(ws, lay.cols["reached"]) or []:
            if self._stage(option) == stage:
                return option
        name = self.stage_names[stage]
        return (self.cfg.statuses.stages[name] or [name])[0]

    def _need(self, *fields: str) -> None:
        missing = [f for f in fields if not self.cfg.has(f)]
        if missing:
            raise TrackerError(f"This needs the column(s) {', '.join(missing)}, "
                               "which the config does not define.")

    # ---- reading

    def version(self) -> str:
        return self.store.version()

    def _layout(self, ws: Worksheet) -> Layout:
        cfg = self.cfg
        header_row = cfg.header_row or next(
            (r for r in range(1, 31) for c in range(1, ws.max_column + 1)
             if str(ws.cell(r, c).value or "").strip() == cfg.columns["company"]), None)
        if header_row is None:
            raise TrackerError(f"Header row not found: no cell «{cfg.columns['company']}» "
                               f"in the first 30 rows of «{cfg.sheet}»")
        headers = {c: str(ws.cell(header_row, c).value).strip()
                   for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value}
        cols = {}
        for key, title in cfg.columns.items():
            found = ([c for c, h in headers.items() if h == title]
                     or [c for c, h in headers.items() if h.startswith(title)])
            if not found:
                raise TrackerError(f"Column «{title}» ({key}) not found in row {header_row}")
            cols[key] = found[0]
        last_row = header_row
        for r in range(header_row + 1, ws.max_row + 1):
            if ws.cell(r, cols["company"]).value not in (None, ""):
                last_row = r
        return Layout(header_row, cols, headers, last_row)

    def _sheet(self, wb: Workbook) -> Worksheet:
        if self.store.single_sheet:
            return wb.worksheets[0]
        name = sheet_in(wb, self.cfg.sheet)
        if name not in wb.sheetnames:
            raise TrackerError(f"Sheet «{self.cfg.sheet}» not found. Sheets: {', '.join(wb.sheetnames)}")
        return wb[name]

    def _load(self) -> tuple[Workbook, Worksheet, Layout, str]:
        """Read once: the workbook and its version come from the same read."""
        wb, version, self.lossy = self.store.load()
        ws = self._sheet(wb)
        return wb, ws, self._layout(ws), version

    def _record(self, ws: Worksheet, lay: Layout, r: int, full: bool = False) -> dict[str, Any]:
        rec: dict[str, Any] = {"row": r}
        for key, c in lay.cols.items():
            rec[key] = self._show(ws.cell(r, c).value, key)
        if full:
            known = set(lay.cols.values())
            other = {h: self._show(ws.cell(r, c).value) for c, h in lay.headers.items()
                     if c not in known and ws.cell(r, c).value not in (None, "")}
            if other:
                rec["other"] = other
        return rec

    def records(self, full: bool = False) -> tuple[list[dict[str, Any]], str]:
        _, ws, lay, version = self._load()
        recs = [self._record(ws, lay, r, full) for r in lay.rows()
                if ws.cell(r, lay.cols["company"]).value not in (None, "")]
        return recs, version

    @staticmethod
    def _label(ws: Worksheet, lay: Layout, r: int) -> str:
        company = str(ws.cell(r, lay.cols["company"]).value or "").strip()
        role = str(ws.cell(r, lay.cols["role"]).value or "").strip()
        return f"{company} — {role}" if role else company

    @staticmethod
    def _brief(rec: dict[str, Any]) -> dict[str, Any]:
        keys = ("row", "num", "company", "role", "status", "applied_on", "followup_on")
        return {k: rec[k] for k in keys if rec.get(k) not in (None, "")}

    def unavailable(self) -> dict[str, str]:
        """Tools that cannot work on this layout, and the column each one misses."""
        need = {"followups_due": ["followup_on"], "log_followup": ["followup_on", "followup_log"],
                "reschedule_followup": ["followup_on"],
                "stale_applications": ["applied_on"], "close_stale": ["applied_on"]}
        out = {}
        for tool, fields in need.items():
            missing = [f for f in fields if not self.cfg.has(f)]
            if missing:
                out[tool] = f"needs the {', '.join(missing)} column (add it to [columns] in the config)"
        return out

    def describe(self) -> dict[str, Any]:
        """How the server reads this tracker: columns, statuses, follow-up rule."""
        _, ws, lay, version = self._load()
        st, fu = self.cfg.statuses, self.cfg.followups
        mapped = set(lay.cols.values())
        out = {
            "tracker": self.store.label,
            "kind": self.store.kind,
            "sheet": ws.title,
            "header_row": lay.header_row,
            "data_rows": f"{lay.header_row + 1}–{lay.last_row}" if lay.last_row > lay.header_row else "none yet",
            "columns": {k: f"{get_column_letter(c)}: {lay.headers[c]}" for k, c in lay.cols.items()},
            "custom_fields": self.cfg.custom_fields(),
            # Not in the config: shown in get_application under "other"; usable by exact
            # header text in add_application(extra=...) and funnel(by=...).
            "other_headers": [h for c, h in lay.headers.items() if c not in mapped],
            "statuses_allowed": self._allowed_statuses(ws, lay),
            "status_not_applied": st.not_applied,
            "status_applied": st.applied,
            "statuses_closed": st.closed,
            "statuses_rejected": st.rejected,
            "funnel_stages": st.stages,
            "followups": {"first_after_days": fu.first_after_days, "next_after_days": fu.next_after_days,
                          "weekend_to_monday": fu.weekend_to_monday},
            **({"date_formats": self.cfg.values.date_formats} if self.cfg.values.date_formats else {}),
            **({"unavailable_tools": u} if (u := self.unavailable()) else {}),
            **({"writes_blocked": self.lossy} if self.lossy else {}),
            "version": version,
        }
        return out

    def list_applications(self, status: str | None = None, query: str | None = None,
                          applied_since: date | None = None, applied_until: date | None = None,
                          open_only: bool = False, limit: int = 100) -> dict[str, Any]:
        recs, version = self.records()
        wanted = None
        if status:
            _, ws, lay, _ = self._load()
            known = {s.casefold(): s for s in self._allowed_statuses(ws, lay)}
            known.update({(r["status"] or "").casefold(): r["status"] for r in recs if r["status"]})
            wanted = set()
            for s in status.split(","):
                s = s.strip().casefold()
                s = self.aliases.get(s, s).casefold()
                if s not in known:
                    raise TrackerError(f"Unknown status «{s}». Statuses in this tracker: {', '.join(known.values())}")
                wanted.add(s)
        q = query.casefold() if query else None
        out = []
        for rec in recs:
            st = (rec["status"] or "").casefold()
            if wanted and st not in wanted:
                continue
            if open_only and rec["status"] in self.closed:
                continue
            if q and q not in f"{rec['company']} {rec['role']}".casefold():
                continue
            if applied_since or applied_until:
                d = self._date(rec.get("applied_on"))
                if (not self._is_applied(rec) or d is None
                        or (applied_since and d < applied_since)
                        or (applied_until and d > applied_until)):
                    continue
            out.append(self._brief(rec))
        return {"count": len(out), "applications": out[:limit],
                "truncated": len(out) > limit, "version": version}

    def get_application(self, query: str | None = None, row: int | None = None) -> dict[str, Any]:
        if query is None and row is None:
            raise TrackerError("Pass a company/role text or a row number")
        _, ws, lay, version = self._load()
        if row is not None:
            self._check_row(ws, lay, row)
            rows = [row]
        else:
            q = query.casefold()
            rows = [r for r in lay.rows()
                    if q in f"{ws.cell(r, lay.cols['company']).value or ''} "
                            f"{ws.cell(r, lay.cols['role']).value or ''}".casefold()]
        return {"count": len(rows), "applications": [self._record(ws, lay, r, full=True) for r in rows],
                "version": version}

    def followups_due(self, within_days: int = 7) -> dict[str, Any]:
        self._need("followup_on")
        recs, version = self.records()
        today = self.today()
        horizon = today + timedelta(days=within_days)
        due, undated = [], []
        for rec in recs:
            if rec["status"] in self.closed or rec["status"] == self.cfg.statuses.not_applied:
                continue
            d = self._date(rec["followup_on"])
            if d is None:
                if self._is_applied(rec) and not rec["followup_on"]:
                    undated.append(self._brief(rec))
                continue
            if d <= horizon:
                log = rec.get("followup_log") or ""
                due.append({**self._brief(rec), "days_left": (d - today).days, "overdue": d < today,
                            "contact": rec.get("followup_contact"),
                            "last_log": log[-300:] if log else None})
        due.sort(key=lambda x: x["followup_on"])
        return {"today": today.isoformat(), "within_days": within_days, "due": due,
                "applied_without_followup_date": undated, "version": version}

    def summary(self, since: date | None = None, until: date | None = None) -> dict[str, Any]:
        recs, version = self.records()
        applied = [r for r in recs if self._is_applied(r)]
        months = Counter()
        undated = 0
        for r in applied:
            d = self._date(r.get("applied_on"))
            if d:
                months[d.strftime("%Y-%m")] += 1
            else:
                undated += 1
        out: dict[str, Any] = {
            "rows": len(recs),
            "applied": len(applied),
            "not_applied": len(recs) - len(applied),
            "by_status": dict(Counter(r["status"] or "—" for r in recs).most_common()),
            "applied_by_month": dict(sorted(months.items())),
            "applied_without_date": undated,
            "version": version,
        }
        if since or until:
            cohort = [r for r in applied if (d := self._date(r.get("applied_on")))
                      and (not since or d >= since) and (not until or d <= until)]
            out["period"] = {
                "since": since.isoformat() if since else None,
                "until": until.isoformat() if until else None,
                "applied": len(cohort),
                "by_status_now": dict(Counter(r["status"] or "—" for r in cohort).most_common()),
                "companies": [f"{r['company']} — {r['role']}" for r in cohort],
            }
        return out

    # ---- funnel

    def _furthest(self, rec: dict[str, Any]) -> int:
        """-1: no reply; 0: a reply (a rejection counts); 1..n: reached stage n."""
        stages = [s for s in (self._stage(rec.get("status")), self._stage(rec.get("reached"))) if s is not None]
        if stages:
            return max(stages) + 1
        return 0 if rec.get("status") in self.rejected else -1

    def _group(self, rec: dict[str, Any], by: str) -> str:
        if by == "month":
            d = self._date(rec.get("applied_on"))
            return d.strftime("%Y-%m") if d else "(no date)"
        text = str(rec.get(by) or "").strip()
        if not text:
            return "(empty)"
        low = text.casefold()
        for name, parts in self.cfg.funnel.by.get(by, self.cfg.funnel.groups).items():
            if any(p.casefold() in low for p in parts):
                return name
        # free text such as "Ashby (direct form), filled in by hand" -> "Ashby"
        head = re.split(r"\s*[(\[,;]|\s+[—–/→|-]\s+", text, maxsplit=1)[0].strip()
        return (head or text)[:40]

    def funnel(self, by: str | None = None, since: date | None = None, until: date | None = None,
               max_groups: int = 12) -> dict[str, Any]:
        """Applied -> replied -> each configured stage, in total and per group."""
        if by is None:  # the channel if the tracker has one, else the month
            by = "channel" if self.cfg.has("channel") else "month" if self.cfg.has("applied_on") else "status"
        other = False
        if by != "month" and not self.cfg.has(by):
            _, _, lay, _ = self._load()
            mapped = set(lay.cols.values())
            headers = [h for c, h in lay.headers.items() if c not in mapped]
            if by not in headers:
                options = ["month", *[f for f in self.cfg.columns if f not in ("company", "role", "num", "link",
                                                                                 "notes", "followup_log")], *headers]
                raise TrackerError(f"Cannot group by «{by}». Options: {', '.join(options)}")
            other = True
        if by == "month" or since or until:
            self._need("applied_on")
        recs, version = self.records(full=other)
        if other:
            for r in recs:
                r[by] = r.get("other", {}).get(by)
        applied = [r for r in recs if self._is_applied(r)]
        skipped = 0
        if since or until:
            kept = []
            for r in applied:
                d = self._date(r.get("applied_on"))
                if d and (not since or d >= since) and (not until or d <= until):
                    kept.append(r)
                elif d is None:
                    skipped += 1
            applied = kept
        levels = ["replied", *self.stage_names]
        no_response = self.cfg.statuses.no_response

        def count(rs: list[dict[str, Any]]) -> dict[str, Any]:
            reached = [self._furthest(r) for r in rs]
            out: dict[str, Any] = {"applied": len(rs)}
            for i, name in enumerate(levels):
                out[name] = sum(1 for x in reached if x >= i)
            out["waiting"] = sum(1 for r in rs if r["status"] in self.waiting)
            out["no_response"] = sum(1 for r in rs if r["status"] == no_response)
            out["rates"] = {name: round(100 * out[name] / len(rs), 1) if rs else None for name in levels}
            return out

        buckets: dict[str, list[dict[str, Any]]] = {}
        names: dict[str, str] = {}
        for r in applied:
            key = self._group(r, by)
            names.setdefault(key.casefold(), key)
            buckets.setdefault(key.casefold(), []).append(r)
        order = sorted(buckets, key=(lambda k: k) if by == "month" else (lambda k: (-len(buckets[k]), k)))
        if by != "month" and len(order) > max_groups:
            merged = order[max_groups - 1:]
            buckets["(other)"] = [r for k in merged for r in buckets[k]]
            names["(other)"] = f"(other: {len(merged)} groups)"
            order = order[:max_groups - 1] + ["(other)"]
        result: dict[str, Any] = {
            "by": by,
            "levels": ["applied", *levels],
            "total": count(applied),
            "groups": [{"group": names[k], **count(buckets[k])} for k in order],
            "note": ("Stage = the later of the current status and the `reached` column."
                     if self.cfg.has("reached") else
                     "Stages come from the current status only, so an application rejected after "
                     "an interview counts as a reply. A `reached` column keeps the furthest stage."),
            "version": version,
        }
        if since or until:
            result["period"] = {"since": since.isoformat() if since else None,
                                "until": until.isoformat() if until else None,
                                "applied_without_date": skipped}
        return result

    def _log_dates(self, text: Any, not_before: date | None = None) -> list[date]:
        """Dates that start entries of a follow-up log ("04.10: InMail sent | 2026-10-18: ...")."""
        if not text:
            return []
        sep = self.cfg.values.separator.strip() or "|"
        msg = self.cfg.messages
        # the server's own "moved" and "dropped" entries say nothing was sent: not a touch
        quiet = [t.split("{")[0].strip().casefold() for t in (msg.followup_moved, msg.followup_dropped)]
        quiet = [q for q in quiet if len(q) >= 4]
        out = []
        for entry in re.split(rf"{re.escape(sep)}|\n", str(text)):
            m = _LEAD_DATE.match(entry)
            if not m:
                continue
            rest = entry[m.end():].lstrip(" :—–-").casefold()
            if any(rest.startswith(q) for q in quiet):
                continue
            if d := date_in_text(m.group(1), self.today(), not_before):
                out.append(d)
        return out

    def _last_touch(self, rec: dict[str, Any]) -> date | None:
        """The submit date, or the newest dated entry of the follow-up log if later."""
        applied = self._date(rec.get("applied_on"))
        if self.cfg.rules.silence_from == "applied":
            return applied
        return max([d for d in (applied, *self._log_dates(rec.get("followup_log"), applied)) if d], default=None)

    def _stale_rows(self, ws: Worksheet, lay: Layout, days: int) -> list[dict[str, Any]]:
        """Still waiting, silent for `days`+ days since the last touch (the submit date or
        the newest dated follow-up log entry), and no follow-up planned ahead."""
        self._need("applied_on")
        today = self.today()
        out = []
        for r in lay.rows():
            if ws.cell(r, lay.cols["company"]).value in (None, ""):
                continue
            rec = self._record(ws, lay, r)
            applied = self._date(rec.get("applied_on"))
            touch = self._last_touch(rec)
            fu = self._date(rec.get("followup_on"))
            if (rec["status"] in self.waiting and self._is_applied(rec) and applied and touch
                    and (today - touch).days >= days and not (fu and fu >= today)):
                log = rec.get("followup_log") or ""
                out.append({**self._brief(rec), "days_since_applied": (today - applied).days,
                            "last_touch": touch.isoformat(), "days_silent": (today - touch).days,
                            "last_log": log[-200:] if log else None})
        return out

    def stale_applications(self, days: int = 30) -> dict[str, Any]:
        _, ws, lay, version = self._load()
        rows = self._stale_rows(ws, lay, days)
        return {"days": days, "count": len(rows), "applications": rows,
                "would_set_status": self.cfg.statuses.no_response, "version": version}

    def close_stale(self, days: int = 30, *, rows: list[int] | None = None,
                    expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        """Set the no-response status on every stale row (or the listed subset) in one write."""
        cfg = self.cfg

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            stale = [x["row"] for x in self._stale_rows(ws, lay, days)]
            if rows is not None:
                extra = sorted(set(rows) - set(stale))
                if extra:
                    raise TrackerError(f"Rows {extra} are not stale by the rule (waiting status, silent "
                                       f"{days}+ days since the last touch, no follow-up planned ahead).")
                stale = [r for r in stale if r in set(rows)]
            if not stale:
                return [], [], []
            status = self._normalize_status(ws, lay, cfg.statuses.no_response)
            text = cfg.messages.stale_closed.format(days=days)
            changes: list[Change] = []
            for r in stale:
                self._set(ws, lay, r, "status", status, changes)
                if cfg.has("followup_on"):
                    self._set(ws, lay, r, "followup_on", cfg.values.no_date, changes)
                self._log(ws, lay, r, text, changes)
            return stale, changes, []

        return self._write("close_stale", mutate, expected_version, dry_run)

    def check_tracker(self, limit: int = 200) -> dict[str, Any]:
        from .doctor import check
        return check(self, limit)

    def fix_tracker(self, kinds: list[str] | None = None, rows: list[int] | None = None, *,
                    expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        from .doctor import fix
        return fix(self, kinds, rows, expected_version=expected_version, dry_run=dry_run)

    def reply_times(self, by: str | None = None, since: date | None = None) -> dict[str, Any]:
        from .insights import reply_times
        return reply_times(self, by, since)

    def followup_events(self, days: int = 60) -> dict[str, Any]:
        from .insights import followup_events
        events, version = followup_events(self, days)
        return {"events": events, "count": len(events), "version": version}

    def export_calendar(self, path: str | Path | None = None, days: int = 60) -> dict[str, Any]:
        from .insights import export_calendar
        return export_calendar(self, path, days)

    def match_emails(self, messages: list[dict[str, Any]]) -> dict[str, Any]:
        from .mail import match
        return match(self, messages)

    def recent_writes(self, limit: int = 10) -> dict[str, Any]:
        entries = self._audit_entries()
        return {"writes": entries[-limit:][::-1]}

    # ---- writing: the tools

    def _company_warnings(self, ws: Worksheet, lay: Layout, company: str, day: date,
                          exclude: int | None = None) -> list[str]:
        """Other submissions to the same company within rules.same_company_days of `day`."""
        n = self.cfg.rules.same_company_days
        key = company_key(company)
        if not n or not key or not self.cfg.has("applied_on"):
            return []
        out = []
        for r in lay.rows():
            if r == exclude or company_key(ws.cell(r, lay.cols["company"]).value) != key:
                continue
            rec = self._record(ws, lay, r)
            d = self._date(rec.get("applied_on"))
            if d and self._is_applied(rec) and abs((day - d).days) < n:
                out.append(f"Row {r} · {self._label(ws, lay, r)}: submitted {d.isoformat()} "
                           f"({abs((day - d).days)} days apart), status {rec.get('status')}. Your rule is one "
                           f"application per company every {n} days (rules.same_company_days).")
        return out

    def add_application(self, company: str, role: str, *, link: str | None = None,
                        notes: str | None = None, extra: dict[str, str] | None = None,
                        allow_duplicate: bool = False, expected_version: str | None = None,
                        dry_run: bool = False) -> dict[str, Any]:
        company, role = company.strip(), role.strip()
        if not company or not role:
            raise TrackerError("company and role are required")
        cfg = self.cfg

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            extra_cols = self._resolve_extra(lay, extra or {})
            if not allow_duplicate:
                for r in lay.rows():
                    same = (str(ws.cell(r, lay.cols["company"]).value or "").strip().casefold() == company.casefold()
                            and str(ws.cell(r, lay.cols["role"]).value or "").strip().casefold() == role.casefold())
                    same_link = (link and cfg.has("link")
                                 and str(ws.cell(r, lay.cols["link"]).value or "").strip() == link.strip())
                    if same or same_link:
                        raise TrackerError(f"Row {r} already has this application. "
                                           "Pass allow_duplicate=true to add anyway.")
            r = lay.last_row + 1
            if lay.last_row > lay.header_row:  # an empty tracker has no data row to copy from
                self._copy_row_style(ws, lay, src=lay.last_row, dst=r)
            values: dict[str, Any] = {"company": company, "role": role, "link": link, "notes": notes}
            if cfg.has("num"):
                nums = [int(v) for x in lay.rows()
                        if str(v := ws.cell(x, lay.cols["num"]).value or "").strip().isdigit()]
                values = {"num": max(nums, default=0) + 1, **values}
            values["status"] = cfg.statuses.not_applied
            if cfg.has("applied"):
                values["applied"] = self._flag(ws, lay, False)
            changes: list[Change] = []
            for key, val in values.items():
                if val not in (None, "") and cfg.has(key):
                    self._set(ws, lay, r, key, val, changes)
            for col, val in extra_cols.items():
                self._set_col(ws, lay, r, col, val, changes)
            layout_notes = self._grow_layout(wb, ws, lay, r)
            return r, changes, layout_notes, self._company_warnings(ws, lay, company, self.today(), exclude=r)

        return self._write("add_application", mutate, expected_version, dry_run)

    def mark_applied(self, row: int, *, applied_on: date | None = None, channel: str | None = None,
                     followup_on: date | None = None, note: str | None = None,
                     expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        cfg, msg = self.cfg, self.cfg.messages

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            self._check_row(ws, lay, row)
            rec = self._record(ws, lay, row)
            if self._is_applied(rec) and (cfg.has("applied") or rec.get("applied_on")):
                raise TrackerError(f"Row {row} is already marked as applied ({rec.get('applied_on')}).")
            day = applied_on or self.today()
            changes: list[Change] = []
            if cfg.has("applied"):
                self._set(ws, lay, row, "applied", self._flag(ws, lay, True), changes)
            if cfg.has("applied_on"):
                self._set(ws, lay, row, "applied_on", day, changes)
            if channel and cfg.has("channel"):
                self._set(ws, lay, row, "channel", channel, changes)
            if ws.cell(row, lay.cols["status"]).value in (None, "", cfg.statuses.not_applied):
                self._set(ws, lay, row, "status", cfg.statuses.applied, changes)
            if cfg.has("followup_on"):
                current_fu = ws.cell(row, lay.cols["followup_on"]).value
                if followup_on or current_fu in (None, ""):
                    fu = followup_on or self._next_date(day, cfg.followups.first_after_days)
                    self._set(ws, lay, row, "followup_on", fu, changes)
                    self._log(ws, lay, row, msg.applied_with_followup.format(date=self._fmt(fu)), changes, on=day)
                else:
                    self._log(ws, lay, row, msg.applied, changes, on=day)
            if note and cfg.has("notes"):
                self._set(ws, lay, row, "notes",
                          self._append(ws.cell(row, lay.cols["notes"]).value, self._stamp(note)), changes)
            warnings = self._company_warnings(ws, lay, rec["company"], day, exclude=row)
            return row, changes, [], warnings

        return self._write("mark_applied", mutate, expected_version, dry_run)

    def set_status(self, row: int, status: str, *, response: str | None = None,
                   note: str | None = None, on: date | None = None, expected_version: str | None = None,
                   dry_run: bool = False) -> dict[str, Any]:
        """`on`: when it happened (the e-mail's date), default today. It goes into the
        reply date when this is the application's first reply."""
        cfg = self.cfg
        if on is not None and on > self.today():
            raise TrackerError(f"{on.isoformat()} is in the future; `on` is when the reply came.")

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            self._check_row(ws, lay, row)
            st = self._normalize_status(ws, lay, status)
            if st == cfg.statuses.applied and not self._is_applied(self._record(ws, lay, row)):
                raise TrackerError("For a new submission use mark_applied: it also sets "
                                   "the date and the follow-up date.")
            changes: list[Change] = []
            had_reply = self._furthest(self._record(ws, lay, row)) >= 0
            self._set(ws, lay, row, "status", st, changes)
            if cfg.has("reached"):  # the furthest stage only moves forward
                cell = ws.cell(row, lay.cols["reached"]).value
                new, old = self._stage(st), self._stage(cell)
                known = old is not None or cell in (None, "")   # never overwrite a word we cannot read
                if new is not None and known and (old is None or new > old):
                    self._set(ws, lay, row, "reached", self._stage_word(ws, lay, new), changes)
            # the first reply only: a row that had one already keeps an empty date for
            # check_tracker to fill from what the row says, not today's date
            if cfg.has("replied_on") and self._is_reply(st) and not had_reply:
                if ws.cell(row, lay.cols["replied_on"]).value in (None, ""):
                    self._set(ws, lay, row, "replied_on", on or self.today(), changes)
            for key, text in (("response", response), ("notes", note)):
                if text and cfg.has(key):
                    self._set(ws, lay, row, key,
                              self._append(ws.cell(row, lay.cols[key]).value, self._stamp(text, on)), changes)
            if cfg.has("followup_on"):
                fu = self._date(ws.cell(row, lay.cols["followup_on"]).value)
                if st in self.closed and fu and fu >= self.today():
                    self._set(ws, lay, row, "followup_on", cfg.values.no_date, changes)
                    self._log(ws, lay, row,
                              cfg.messages.followup_cancelled.format(date=self._fmt(fu), status=st), changes)
            return row, changes, []

        return self._write("set_status", mutate, expected_version, dry_run)

    def log_followup(self, row: int, text: str, *, next_on: date | None = None, final: bool = False,
                     contact: str | None = None, expected_version: str | None = None,
                     dry_run: bool = False) -> dict[str, Any]:
        self._need("followup_on", "followup_log")
        cfg = self.cfg

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            self._check_row(ws, lay, row)
            changes: list[Change] = []
            self._log(ws, lay, row, text, changes)
            nxt = cfg.values.no_date if final else (
                next_on or self._next_date(self.today(), cfg.followups.next_after_days))
            self._set(ws, lay, row, "followup_on", nxt, changes)
            if contact and cfg.has("followup_contact"):
                self._set(ws, lay, row, "followup_contact", contact, changes)
            return row, changes, []

        return self._write("log_followup", mutate, expected_version, dry_run)

    def reschedule_followup(self, row: int, on: date | None, reason: str, *,
                            expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        """Move the next follow-up to another date, or plan none (on=None) while the
        application stays open: nobody to write to yet, waiting for LinkedIn credits,
        the company promised an answer by a date. Nothing was sent, so nothing is
        counted as a follow-up; the reason goes to the follow-up log."""
        self._need("followup_on")
        cfg = self.cfg
        if not reason.strip():
            raise TrackerError("Say why the follow-up moves; the reason goes to the follow-up log.")
        if on is not None and on < self.today():
            raise TrackerError(f"{on.isoformat()} is in the past; pick today or a later date.")

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            self._check_row(ws, lay, row)
            status = ws.cell(row, lay.cols["status"]).value
            if status in self.closed:
                raise TrackerError(f"Row {row} is closed ({status}); it has no follow-ups to move.")
            changes: list[Change] = []
            self._set(ws, lay, row, "followup_on", on if on is not None else cfg.values.no_date, changes)
            text = (cfg.messages.followup_moved.format(date=self._fmt(on), reason=reason.strip()) if on
                    else cfg.messages.followup_dropped.format(reason=reason.strip()))
            self._log(ws, lay, row, text, changes)
            return row, changes, []

        return self._write("reschedule_followup", mutate, expected_version, dry_run)

    # fields edit_application leaves to the tools that know their side effects
    _EDIT_ELSEWHERE = {
        "status": "set_status (it also moves the furthest stage, fills the reply date and cancels a "
                  "pending follow-up)",
    }

    def edit_application(self, row: int, fields: dict[str, Any], reason: str, *,
                         expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        """Correct fields of one row by hand: a wrong date, a typo, the furthest stage, a
        custom column. `fields`: field name -> new value (null clears the cell). Dates are
        checked and written as dates; the reason, with what changed, goes to the notes."""
        cfg = self.cfg
        if not reason.strip():
            raise TrackerError("Say why the row is corrected; the reason goes to the notes.")
        if not fields:
            raise TrackerError("Name at least one field to change.")
        unknown = [f for f in fields if f not in cfg.columns]
        if unknown:
            raise TrackerError(f"Unknown field(s) {', '.join(unknown)}. Fields of this tracker: "
                               + ", ".join(f for f in cfg.columns if f not in self._EDIT_ELSEWHERE))
        for f in fields:
            if f in self._EDIT_ELSEWHERE:
                raise TrackerError(f"Change {f} with {self._EDIT_ELSEWHERE[f]}.")
        today = self.today()
        new: dict[str, Any] = {}
        for f, v in fields.items():
            v = v.strip() if isinstance(v, str) else v
            if v in (None, ""):
                if f == "company":
                    raise TrackerError("The company cannot be empty: it is how a row is found.")
                new[f] = None
            elif cfg.is_date(f):
                if str(v) == cfg.values.no_date:
                    new[f] = cfg.values.no_date
                    continue
                d = self._date(v)
                if d is None:
                    raise TrackerError(f"{f}: «{v}» is not a date (use YYYY-MM-DD).")
                if f in ("applied_on", "replied_on") and d > today:
                    raise TrackerError(f"{f}: {d.isoformat()} is in the future.")
                new[f] = d
            elif f == "applied":
                yes = str(v).casefold() in (cfg.values.yes.casefold(), "true", "yes", "да")
                no = str(v).casefold() in (cfg.values.no.casefold(), "false", "no", "нет")
                if not (yes or no):
                    raise TrackerError(f"applied: «{v}» — use «{cfg.values.yes}» or «{cfg.values.no}».")
                new[f] = cfg.values.yes if yes else cfg.values.no
            elif f == "reached":
                if self._stage(v) is None:
                    raise TrackerError(f"reached: «{v}» is not a stage. Stages: "
                                       + ", ".join(w for words in cfg.statuses.stages.values() for w in words))
                new[f] = v   # written as the column's own word below
            else:
                new[f] = str(v)

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            self._check_row(ws, lay, row)
            missing = [f for f in new if f not in lay.cols]
            if missing:
                raise TrackerError(f"Not in the sheet: {', '.join(missing)} (see describe_tracker).")
            rec = self._record(ws, lay, row)
            applied = new.get("applied_on", self._date(rec.get("applied_on")))
            replied = new.get("replied_on", self._date(rec.get("replied_on")))
            if isinstance(applied, date) and isinstance(replied, date) and replied < applied:
                raise TrackerError(f"The reply date {replied.isoformat()} would be before the submit date "
                                   f"{applied.isoformat()}.")
            warnings = []
            changes: list[Change] = []
            for f, v in new.items():
                if f == "reached" and v is not None:
                    v = self._stage_word(ws, lay, self._stage(v))
                    old = self._stage(ws.cell(row, lay.cols["reached"]).value)
                    if old is not None and self._stage(v) < old:
                        warnings.append(f"The furthest stage goes back from «{rec.get('reached')}» to «{v}».")
                if f == "followup_on" and isinstance(v, date):
                    if v.weekday() >= 5 and cfg.followups.weekend_to_monday:
                        warnings.append(f"{v.isoformat()} is a {v.strftime('%A')}; the follow-up rule "
                                        "would use Monday.")
                    if rec["status"] in self.closed:
                        warnings.append(f"The application is closed ({rec['status']}).")
                self._set(ws, lay, row, f, v, changes)
            if changes and cfg.has("notes") and "notes" not in new:
                def shown(v: Any) -> Any:
                    return "∅" if v in (None, "") else v
                what = "; ".join(f"{lay.headers.get(lay.cols[c.field], c.field)} «{shown(c.old)}» "
                                 f"→ {shown(self._show(c.new, c.field))}" for c in changes)
                self._set(ws, lay, row, "notes", self._append(ws.cell(row, lay.cols["notes"]).value,
                                                              self._stamp(f"{what} — {reason.strip()}")), changes)
            return row, changes, [], warnings

        return self._write("edit_application", mutate, expected_version, dry_run)

    def add_column(self, header: str, *, field: str | None = None, kind: str = "text",
                   expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        """A new column after the last one, with the header styled like its neighbour.
        `field` also maps it in the config ([columns] field = "header"), so the tools use
        it: replied_on for reply dates, reached for the furthest stage, or a custom name.
        `kind`: "date" formats the cells as dates (implied by a date field). Undo removes
        the column and the config line together."""
        cfg = self.cfg
        header = (header or "").strip()
        if not header or len(header) > 100 or "\n" in header:
            raise TrackerError("The header must be one line of 1–100 characters.")
        if kind not in ("text", "date"):
            raise TrackerError("kind is «text» or «date».")
        new_text = None
        if field is not None:
            if not re.fullmatch(r"[a-z][a-z0-9_]{0,40}", field):
                raise TrackerError("A field name is lowercase letters, digits and _, e.g. replied_on or source.")
            if field in ("company", "role", "status"):
                raise TrackerError(f"{field} is a column every tracker has already.")
            if field in cfg.columns:
                raise TrackerError(f"{field} is already the column «{cfg.columns[field]}».")
            if self.config_file is None or not self.config_file.exists():
                raise TrackerError("The server does not know where this tracker's config is, so it cannot map "
                                   f"a field. Add the column without `field`, then add {field} = \"{header}\" "
                                   "to [columns] in the config by hand.")
            new_text = with_column(self.config_file.read_text(encoding="utf-8"), field, header)
            parse_config(new_text, str(self.config_file))   # refuse before writing anything
            if cfg.is_date(field) or field.endswith("_on"):
                kind = "date"
        config_line = f"{field} = {json.dumps(header, ensure_ascii=False)}" if field else None

        def mutate(wb: Workbook, ws: Worksheet, lay: Layout):
            if any(h.casefold() == header.casefold() for h in lay.headers.values()):
                raise TrackerError(f"There is a column «{header}» already"
                                   + (f"; to use it, add {config_line} to [columns] in the config." if field else "."))
            last = max(lay.headers)
            col = last + 1
            if any(ws.cell(r, col).value not in (None, "") for r in range(lay.header_row, ws.max_row + 1)):
                raise TrackerError(f"Column {get_column_letter(col)} has values but no header; give it one or "
                                   "clear it first.")
            prev, cell = ws.cell(lay.header_row, last), ws.cell(lay.header_row, col)
            if isinstance(cell, MergedCell):
                raise TrackerError(f"{cell.coordinate} is part of a merged range; unmerge it first.")
            changes: list[Change] = []
            self._set_cell(cell, lay.header_row, header, header, changes)
            notes = [f"column {get_column_letter(col)} «{header}» after «{lay.headers[last]}»"]
            if self.store.kind == "xlsx":
                cell._style = copy(prev._style)
                letter, prev_letter = get_column_letter(col), get_column_letter(last)
                if prev_letter in ws.column_dimensions and ws.column_dimensions[prev_letter].width:
                    ws.column_dimensions[letter].width = ws.column_dimensions[prev_letter].width
                date_fmt = None
                if kind == "date":
                    # the submit dates' most common format, without a time of day
                    src = lay.cols.get("applied_on") or lay.cols.get("followup_on")
                    used = Counter(ws.cell(r, src).number_format for r in range(lay.header_row + 1, lay.last_row + 1)
                                   if src and isinstance(ws.cell(r, src).value, datetime)
                                   and ws.cell(r, src).number_format not in ("General", None)
                                   and "h" not in ws.cell(r, src).number_format.lower())
                    date_fmt = used.most_common(1)[0][0] if used else "yyyy-mm-dd"
                for r in range(lay.header_row + 1, max(lay.last_row, lay.header_row + 1) + 1):
                    target = ws.cell(r, col)
                    target._style = copy(ws.cell(r, last)._style)
                    if date_fmt:
                        target.number_format = date_fmt
                ref = ws.auto_filter.ref
                m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", ref or "")
                if m and m.group(3) == prev_letter:
                    ws.auto_filter.ref = f"{m.group(1)}{m.group(2)}:{letter}{m.group(4)}"
                    notes.append(f"filter now {ws.auto_filter.ref}")
            return lay.header_row, changes, notes

        def write_config() -> dict[str, Any]:
            if not new_text:
                return {}
            assert self.config_file is not None
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            backup = backup_path(self.backup_dir, self.config_file.name + ".before-add-column", "")
            shutil.copy2(self.config_file, backup)
            self.config_file.write_text(new_text, encoding="utf-8")
            self._use(parse_config(new_text, str(self.config_file)))
            return {"config_file": str(self.config_file), "config_backup": backup.name, "config_line": config_line}

        result = self._write("add_column", mutate, expected_version, dry_run, after_save=write_config)
        result["column"] = {"header": header, "kind": kind, **({"field": field, "config_line": config_line}
                                                                if field else {})}
        return result

    def undo_last_write(self, expected_version: str | None = None, dry_run: bool = False) -> dict[str, Any]:
        """Restore the backup of the last write, only if nothing changed since."""
        entries = self._audit_entries()
        entry = next((e for e in reversed(entries) if e["tool"] != "undo_last_write"), None)
        if entry is None:
            raise TrackerError("No writes to undo")
        if entries[-1]["tool"] == "undo_last_write" and undid(entries[-1], entry):
            raise TrackerError("The last write is already undone")
        if entry.get("unverified"):
            raise TrackerError("The last write could not be checked after it was sent, so undo cannot tell "
                               f"what to restore safely. The old values are in {entry['backup']}.")
        current = self.version()
        if current != entry["version_after"] or (expected_version and expected_version != current):
            raise ConflictError("The tracker changed after that write; undo would lose those changes. "
                                f"Restore by hand from {entry['backup']} if needed.")
        if dry_run:
            return {"dry_run": True, "would_undo": entry, "version": current}
        safety, after = self.store.restore(entry["backup"])
        if entry.get("config_backup") and entry.get("config_file"):   # add_column also mapped a field
            source = self.backup_dir / entry["config_backup"]
            if source.exists():
                Path(entry["config_file"]).write_bytes(source.read_bytes())
                self._use(load_config_file(Path(entry["config_file"])))
        self._audit({"tool": "undo_last_write", "undid": entry["at"], "undid_backup": entry.get("backup"),
                     "row": entry["row"],
                     "changes": [], "backup": safety,
                     "version_before": current, "version_after": after})
        return {"undone": entry, "backup_of_state_before_undo": str(self.backup_dir / safety), "version": after}

    # ---- writing: the safe path

    def _write(self, tool: str, mutate: Callable[[Workbook, Worksheet, Layout], tuple],
               expected_version: str | None, dry_run: bool,
               after_save: Callable[[], dict[str, Any]] | None = None) -> dict[str, Any]:
        if not dry_run:
            self.store.check_writable()
        wb, ws, lay, before = self._load()
        if self.lossy and not dry_run:
            raise TrackerError("Writes are off for this workbook: openpyxl, which the server writes with, "
                               "cannot keep some of its Excel features and would drop them on save: "
                               f"{'; '.join(self.lossy)}. A common cause is a dropdown that takes its list "
                               "from another sheet, saved by newer Excel. Reading works as usual.")
        if expected_version and expected_version != before:
            raise ConflictError(f"The tracker changed since you read it (now {before}, "
                                f"you had {expected_version}). Re-read and retry.")
        row, changes, layout_notes, *more = mutate(wb, ws, lay)
        result: dict[str, Any] = {"row": row, "dry_run": dry_run,
                                  "changes": [c.as_dict(self._show) for c in changes],
                                  "labels": {r: self._label(ws, lay, r) for r in sorted({c.row for c in changes})}}
        if layout_notes:
            result["layout"] = layout_notes
        if more and more[0]:   # warnings: shown with the preview, they do not stop the write
            result["warnings"] = more[0]
        if self.lossy:
            result["warning"] = f"This workbook cannot be written: {'; '.join(self.lossy)}"
        if dry_run or not changes:
            result["version"] = before
            if not changes:
                result["note"] = "Nothing to change"
            return result

        try:
            backup, after = self.store.save(wb, before, lambda saved: self._verify(saved, changes))
        except UnverifiedWrite as e:   # Google Sheets: sent, but not read back
            self._audit({"tool": tool, "row": row, "changes": result["changes"], "layout": layout_notes,
                         "backup": e.backup, "version_before": before, "version_after": None,
                         "unverified": str(e)})
            raise
        extra = after_save() if after_save else {}   # e.g. the config line of a new column
        self._audit({"tool": tool, "row": row, "changes": result["changes"], "labels": result["labels"],
                     "layout": layout_notes, "backup": backup,
                     "version_before": before, "version_after": after, **extra})
        result.update(backup=str(self.backup_dir / backup), version=after, **extra)
        return result

    def _verify(self, saved: Workbook, changes: list[Change]) -> None:
        ws = self._sheet(saved)
        lay = self._layout(ws)
        for ch in changes:
            col = lay.cols[ch.field] if ch.field in lay.cols else self._column_of(lay, ch.field)
            got = ws.cell(ch.row, col).value
            if isinstance(ch.new, date):
                ok = self._date(got) == ch.new
            else:  # a CSV gives everything back as text
                ok = got == ch.new or (got is not None and str(got) == str(ch.new))
            if not ok:
                raise TrackerError(f"Verification failed for row {ch.row} {ch.field}: "
                                   f"wrote {ch.new!r}, read back {got!r}. Nothing was saved.")

    def _audit(self, entry: dict[str, Any]) -> None:
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        entry = {"at": datetime.now().isoformat(timespec="seconds"), **entry}
        with open(self.backup_dir / "audit.jsonl", "a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False, default=str) + "\n")

    def _audit_entries(self) -> list[dict[str, Any]]:
        log = self.backup_dir / "audit.jsonl"
        if not log.exists():
            return []
        return [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()]

    # ---- writing: helpers

    def _log(self, ws: Worksheet, lay: Layout, row: int, text: str, changes: list[Change],
             on: date | None = None) -> None:
        if not self.cfg.has("followup_log"):
            return
        old = ws.cell(row, lay.cols["followup_log"]).value
        self._set(ws, lay, row, "followup_log", self._append(old, self._stamp(text, on)), changes)

    def _set(self, ws: Worksheet, lay: Layout, row: int, key: str, new: Any, changes: list[Change]) -> None:
        self._set_cell(ws.cell(row, lay.cols[key]), row, key, new, changes)

    def _set_col(self, ws: Worksheet, lay: Layout, row: int, col: int, new: Any, changes: list[Change]) -> None:
        key = next((k for k, c in lay.cols.items() if c == col), lay.headers[col])
        self._set_cell(ws.cell(row, col), row, key, new, changes)

    def _set_cell(self, cell, row: int, key: str, new: Any, changes: list[Change]) -> None:
        if isinstance(cell, MergedCell):
            raise TrackerError(f"Row {row}, {key}: the cell is part of a merged range; unmerge it in the "
                               "spreadsheet to write there.")
        old = cell.value
        if isinstance(new, date):
            if self._date(old) == new and isinstance(old, datetime):
                return
            cell.value = datetime(new.year, new.month, new.day)
            cell.number_format = self._date_format(cell)
            if old not in (None, "") and not isinstance(old, (datetime, date)):
                # text or a number becoming a date: show what was there, not its reading
                changes.append(Change(row, key, old, new))
                return
        else:
            new = self._as_text(new)
            if old == new:
                return
            cell.value = new
            if isinstance(new, str) and new.startswith("="):
                cell.data_type = "s"   # text from the model stays text, never a formula
        changes.append(Change(row, key, self._show(old, key), new))

    @staticmethod
    def _date_format(cell) -> str:
        """The number format of the column's other dates, so a fixed date looks like them."""
        cells = getattr(cell.parent, "_cells", {})   # existing cells only: reading must not create any
        for (r, c), other in cells.items():
            if c == cell.column and r != cell.row and isinstance(other.value, datetime) \
                    and other.number_format not in ("General", None):
                return other.number_format
        return "yyyy-mm-dd"

    def _as_text(self, value: Any) -> Any:
        """A CSV file has no cell types: Excel runs anything that looks like a formula
        when it opens one. Text the server writes there gets a leading apostrophe if
        it starts like a formula ("=", "+", "@", or "-" followed by more than a space)."""
        if (self.store.kind == "csv" and isinstance(value, str) and value[:1] in ("=", "+", "@", "-")
                and not (value[:1] == "-" and value[1:2] in ("", " "))):
            return "'" + value
        return value

    @staticmethod
    def _column_of(lay: Layout, header: str) -> int:
        return next(c for c, h in lay.headers.items() if h == header)

    def _resolve_extra(self, lay: Layout, extra: dict[str, str]) -> dict[int, Any]:
        """add_application's `extra`: keys are config field names or header texts.
        Columns the server manages are refused whichever way they are named, and
        values for date columns must be dates."""
        out: dict[int, Any] = {}
        reserved = {"num", "company", "role", "status", "applied", "applied_on", "reached",
                    "replied_on", "followup_on", "followup_log"}
        by_col = {c: k for k, c in lay.cols.items()}
        for key, val in extra.items():
            col = lay.cols.get(key) or next((c for c, h in lay.headers.items() if h == key), None)
            if col is None:
                raise TrackerError(f"Unknown column «{key}». Use a field from describe_tracker "
                                   "or an exact header text.")
            field = by_col.get(col, key)
            if field in reserved:
                raise TrackerError(f"«{key}» is set by the server ({field}): use add_application's own "
                                   "arguments, then mark_applied, set_status or log_followup.")
            if self.cfg.is_date(field):
                d = self._date(val)
                if d is None:
                    raise TrackerError(f"«{val}» is not a date for «{key}»; use yyyy-mm-dd.")
                val = d
            out[col] = val
        return out

    def _check_row(self, ws: Worksheet, lay: Layout, row: int) -> None:
        if not (lay.header_row < row <= lay.last_row) or ws.cell(row, lay.cols["company"]).value in (None, ""):
            raise TrackerError(f"Row {row} is not an application row "
                               f"(data rows are {lay.header_row + 1}–{lay.last_row}).")

    def _normalize_status(self, ws: Worksheet, lay: Layout, value: str) -> str:
        allowed = self._allowed_statuses(ws, lay)
        v = value.strip()
        by_fold = {s.casefold(): s for s in allowed}
        if v.casefold() in by_fold:
            return by_fold[v.casefold()]
        alias = self.aliases.get(v.casefold())
        if alias in allowed:
            return alias
        where = (f"the dropdown in column {get_column_letter(lay.cols['status'])}"
                 if dropdown_values(ws, lay.cols["status"])
                 else "the config's statuses.fallback; the column has no dropdown")
        raise TrackerError(f"Unknown status «{value}». Allowed ({where}): {', '.join(allowed)}")

    def _allowed_statuses(self, ws: Worksheet, lay: Layout) -> list[str]:
        return dropdown_values(ws, lay.cols["status"]) or list(self.cfg.statuses.fallback)

    @staticmethod
    def _copy_row_style(ws: Worksheet, lay: Layout, src: int, dst: int) -> None:
        for c in lay.headers:
            a, b = ws.cell(src, c), ws.cell(dst, c)
            if a.has_style:
                b.font, b.fill, b.border = copy(a.font), copy(a.fill), copy(a.border)
                b.alignment, b.protection = copy(a.alignment), copy(a.protection)
                b.number_format = a.number_format
        if ws.row_dimensions[src].height:
            ws.row_dimensions[dst].height = ws.row_dimensions[src].height

    def _grow_layout(self, wb: Workbook, ws: Worksheet, lay: Layout, new_row: int) -> list[str]:
        """Stretch dropdowns, the filter and the summary sheet's formulas to cover new_row."""
        notes = []
        # Only cell values reach a Google Sheet, so its dropdowns are not stretched here
        # (Sheets dropdowns usually cover the whole column anyway).
        for dv in ws.data_validations.dataValidation if self.store.kind != "gsheet" else []:
            ranges = list(dv.sqref.ranges)
            grown = False
            for rng in ranges:
                if rng.min_row <= lay.header_row + 1 and rng.max_row < new_row:
                    rng.expand(down=new_row - rng.max_row)
                    grown = True
            if grown:
                dv.sqref = MultiCellRange(" ".join(r.coord for r in ranges))
                notes.append(f"dropdown {dv.sqref} now reaches row {new_row}")
        if ws.auto_filter.ref:
            m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", ws.auto_filter.ref)
            if m and int(m.group(4)) < new_row:
                ws.auto_filter.ref = f"{m.group(1)}{m.group(2)}:{m.group(3)}{new_row}"
                notes.append(f"filter now {ws.auto_filter.ref}")
        summary = self.cfg.summary_sheet and sheet_in(wb, self.cfg.summary_sheet)
        if summary and summary in wb.sheetnames:
            sheet = self.cfg.sheet
            ref = re.escape(sheet) if re.fullmatch(r"\w+", sheet) else re.escape(f"'{sheet}'")
            pattern = re.compile(rf"({ref}!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?)(\d+)")
            grown_cells = []
            for row in wb[summary].iter_rows():
                for cell in row:
                    f = cell.value
                    if isinstance(f, str) and f.startswith("=") and cell.data_type == "f":
                        new_f = pattern.sub(lambda m: m.group(1) + str(max(int(m.group(2)), new_row + 100))
                                            if int(m.group(2)) < new_row else m.group(0), f)
                        if new_f != f:
                            cell.value = new_f
                            grown_cells.append(cell.coordinate)
            if grown_cells:
                notes.append(f"«{summary}» formulas extended to row {new_row + 100}: "
                             f"{', '.join(grown_cells)}")
        return notes
